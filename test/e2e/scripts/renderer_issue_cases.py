"""Small reproducible native PowerPoint fixtures for renderer issues #36 and #37."""
from copy import deepcopy

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def _label(shapes, text, x, y, w=3.0, h=0.8, preset=MSO_SHAPE.RECTANGLE,
           anchor=MSO_VERTICAL_ANCHOR.MIDDLE):
    shape = shapes.add_shape(preset, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xEF, 0xF5, 0xFB)
    shape.line.color.rgb = RGBColor(0x44, 0x72, 0xC4)
    tf = shape.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(0x10, 0x10, 0x10)
    return shape


def _hidden(shape, value):
    shape._element.xpath(".//p:cNvPr")[0].set("hidden", value)
    return shape


def _build_hidden_nodes(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, value in enumerate([None, "0", "false"]):
        _label(slide.shapes, f"VISIBLE {value or 'absent'}", 0.8 + index * 4, 0.6)
    for index, value in enumerate(["1", "true"]):
        _hidden(_label(slide.shapes, f"HIDDEN shape {value}", 0.8 + index * 4, 1.8), value)
        group = slide.shapes.add_group_shape()
        _label(group.shapes, f"HIDDEN group {value}", 0.8 + index * 4, 3.0)
        _hidden(group, value)
    _label(slide.shapes, "VISIBLE sibling", 8.8, 3.0)
    for index, template in enumerate([slide.slide_layout, slide.slide_layout.slide_master]):
        source = _hidden(_label(slide.shapes, "HIDDEN template", 0.8 + index * 4, 4.8), "true")
        element = deepcopy(source._element)
        element.xpath(".//p:cNvPr")[0].set("id", str(5000 + index))
        template._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld").find("{http://schemas.openxmlformats.org/presentationml/2006/main}spTree").insert_element_before(element, "p:extLst")
        source._element.getparent().remove(source._element)


def _build_preset_text_rectangles(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _label(slide.shapes, "Sample text", 0.8, 0.7, h=1.5, preset=MSO_SHAPE.FLOWCHART_DOCUMENT)
    _label(slide.shapes, "Sample text", 4.8, 0.7, h=2.5, preset=MSO_SHAPE.FLOWCHART_DOCUMENT)
    _label(slide.shapes, "Sample text", 8.8, 0.7, h=1.5)
    for index, anchor in enumerate([MSO_VERTICAL_ANCHOR.TOP, MSO_VERTICAL_ANCHOR.MIDDLE, MSO_VERTICAL_ANCHOR.BOTTOM]):
        _label(slide.shapes, "Anchor text", 0.8 + index * 4, 4.5, h=1.8,
               preset=MSO_SHAPE.FLOWCHART_DOCUMENT, anchor=anchor)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, preset in enumerate([MSO_SHAPE.FLOWCHART_MULTIDOCUMENT, MSO_SHAPE.WAVE, MSO_SHAPE.CLOUD]):
        _label(slide.shapes, "Sample text", 0.8 + index * 4, 1.0, h=2.5, preset=preset)
    group = slide.shapes.add_group_shape()
    _label(group.shapes, "Grouped text", 0.8, 4.5, h=1.5, preset=MSO_SHAPE.FLOWCHART_DOCUMENT)
    group.width = Inches(4.5)
    group.height = Inches(2.0)


def issue_case_definitions():
    return [
        {
            "name": "oracle-pypptx-composite-0021-hidden-nodes",
            "build_fn": _build_hidden_nodes,
            "coverage": {
                "oracle": "native-powerpoint",
                "requiredFonts": ["Arial"],
                "features": ["node.hidden.true-false", "node.hidden.group-subtree", "node.hidden.layout-master"],
            },
        },
        {
            "name": "oracle-pypptx-text-0064-preset-text-rectangles",
            "build_fn": _build_preset_text_rectangles,
            "slide_count": 2,
            "coverage": {
                "oracle": "native-powerpoint",
                "requiredFonts": ["Arial"],
                "features": ["text.preset-rectangle", "text.anchor.top-center-bottom", "text.preset-rectangle.group-scale"],
            },
        },
    ]
