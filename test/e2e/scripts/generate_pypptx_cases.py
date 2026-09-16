#!/usr/bin/env python3
"""Generate ground truth cases using python-pptx + native PowerPoint export.

Produces oracle-pypptx-* cases covering:
  - Rich text: fonts, sizes, bold/italic, alignment, vertical text, bullets
  - Shape adjustment variants: same shape with different adj values
  - Zero-adjustment flowchart geometry: 28 presets across square, wide, and grouped tall views
  - Static DrawingML 3D: bounded orthographic top-bevel and opt-out matrix
  - Tables: isolated style, merge, sizing, alignment, text, and border interactions
  - Formulas: DrawingML a14:m / OMML constructs with explicit MCE shape fallbacks
  - Chart data variants: 2D chart types with custom data/series
  - Composite: multiple components on a single slide

Usage (from test/e2e/):
  pip install python-pptx
  python scripts/generate_pypptx_cases.py              # PDF on macOS; PDF+PNG on Windows
  python scripts/generate_pypptx_cases.py --pptx-only  # generate PPTX only (any platform)
  python scripts/generate_pypptx_cases.py --case 'oracle-pypptx-text-00[456]*'
  python scripts/generate_pypptx_cases.py --include-local-shape3d-matrix \
    --case 'oracle-local-shape3d-*'                     # ignored discovery corpus
"""
from __future__ import annotations

import argparse
import fnmatch
import hashlib
import json
import math
import posixpath
import random
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

E2E_DIR = Path(__file__).resolve().parents[1]
if str(E2E_DIR) not in sys.path:
    sys.path.insert(0, str(E2E_DIR))

CASES_DIR = E2E_DIR / "oracle" / "cases-pypptx"
TESTDATA_DIR = E2E_DIR / "testdata"
REPORT_PATH = E2E_DIR / "reports" / "oracle-failures" / "pypptx-ground-truth.json"
LOCAL_SHAPE3D_CASES_DIR = E2E_DIR / "oracle-runtime" / "local-shape3d-cases"

# Slide dimensions (standard widescreen 13.333" x 7.5")
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---------------------------------------------------------------------------
# Case definition helpers
# ---------------------------------------------------------------------------

CaseDef = dict  # {name: str, build_fn: Callable[[Presentation], None]}

PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PML_REL_PREFIX = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"p": PML_NS, "pr": REL_NS}
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _emu(inches: float) -> int:
    return int(Inches(inches))


def _remove_children(parent, local_names: tuple[str, ...]) -> None:
    for local_name in local_names:
        child = parent.find(qn(f"a:{local_name}"))
        if child is not None:
            parent.remove(child)


def _configure_text_body(
    text_frame,
    *,
    wrap: str,
    autofit: str | None,
    autofit_attrs: dict[str, str] | None = None,
    anchor: str | None = None,
    horz_overflow: str | None = None,
    vert_overflow: str | None = None,
) -> None:
    """Set bodyPr values that python-pptx does not expose losslessly."""
    body_pr = text_frame._txBody.find(qn("a:bodyPr"))
    if body_pr is None:
        raise RuntimeError("text frame has no a:bodyPr")

    body_pr.set("wrap", wrap)
    if anchor is not None:
        body_pr.set("anchor", anchor)
    for attr_name, attr_value in (
        ("horzOverflow", horz_overflow),
        ("vertOverflow", vert_overflow),
    ):
        if attr_value is None:
            body_pr.attrib.pop(attr_name, None)
        else:
            body_pr.set(attr_name, attr_value)
    _remove_children(body_pr, ("noAutofit", "normAutofit", "spAutoFit"))
    if autofit is not None:
        etree.SubElement(body_pr, qn(f"a:{autofit}"), **(autofit_attrs or {}))


def _replace_spacing_value(paragraph, container_name: str, value_name: str, value: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn(f"a:{container_name}"))
    if old is not None:
        p_pr.remove(old)
    container = etree.SubElement(p_pr, qn(f"a:{container_name}"))
    etree.SubElement(container, qn(f"a:{value_name}"), val=str(value))


def _set_cjk_run_style(
    run,
    *,
    font_name: str = "Microsoft YaHei",
    font_size_pt: int = 28,
    bold: bool = False,
    spacing: int | None = None,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    r_pr = run._r.get_or_add_rPr()
    r_pr.set("lang", "zh-CN")
    if spacing is not None:
        r_pr.set("spc", str(spacing))
    for script in ("latin", "ea"):
        typeface = r_pr.find(qn(f"a:{script}"))
        if typeface is None:
            typeface = etree.SubElement(r_pr, qn(f"a:{script}"))
        typeface.set("typeface", font_name)


def _add_cjk_run(paragraph, text: str, **style):
    run = paragraph.add_run()
    run.text = text
    _set_cjk_run_style(run, **style)
    return run


def _replace_text_solid_fill(parent, color_kind: str, color_value: str) -> None:
    """Write a schema-ordered DrawingML text color without python-pptx defaults."""
    for fill_name in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
        old = parent.find(qn(f"a:{fill_name}"))
        if old is not None:
            parent.remove(old)
    solid_fill = etree.Element(qn("a:solidFill"))
    etree.SubElement(solid_fill, qn(f"a:{color_kind}"), val=color_value)
    parent.insert(0, solid_fill)


def _set_shape_font_ref_color(shape, scheme_color: str) -> None:
    """Add a complete p:style tuple with fontRef as the observable fallback color."""
    shape_element = shape._element
    old_style = shape_element.find(qn("p:style"))
    if old_style is not None:
        shape_element.remove(old_style)
    style = etree.Element(qn("p:style"))
    for tag_name, index in (("lnRef", "0"), ("fillRef", "0"), ("effectRef", "0")):
        reference = etree.SubElement(style, qn(f"a:{tag_name}"), idx=index)
        etree.SubElement(reference, qn("a:schemeClr"), val="accent1")
    font_ref = etree.SubElement(style, qn("a:fontRef"), idx="minor")
    etree.SubElement(font_ref, qn("a:schemeClr"), val=scheme_color)
    tx_body = shape_element.find(qn("p:txBody"))
    if tx_body is None:
        raise RuntimeError("shape has no p:txBody")
    shape_element.insert(shape_element.index(tx_body), style)


def _set_paragraph_default_color(paragraph, color_kind: str, color_value: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    old_default = p_pr.find(qn("a:defRPr"))
    if old_default is not None:
        p_pr.remove(old_default)
    default_run_properties = etree.SubElement(p_pr, qn("a:defRPr"))
    _replace_text_solid_fill(default_run_properties, color_kind, color_value)


def _set_run_explicit_color(run, color_kind: str, color_value: str) -> None:
    _replace_text_solid_fill(run._r.get_or_add_rPr(), color_kind, color_value)


def _set_cjk_paragraph_text(paragraph, text: str, **style) -> None:
    """Set paragraph text so vertical tabs become OOXML soft line breaks."""
    paragraph.text = text
    for run in paragraph.runs:
        _set_cjk_run_style(run, **style)


def _shape3d_fixture_image() -> BytesIO:
    """Return a deterministic image whose edge lighting remains visible in native exports."""
    image = Image.new("RGB", (480, 280), "#17456B")
    draw = ImageDraw.Draw(image)
    for x in range(image.width):
        red = 24 + round(50 * x / max(1, image.width - 1))
        green = 74 + round(90 * x / max(1, image.width - 1))
        blue = 112 + round(65 * x / max(1, image.width - 1))
        draw.line((x, 0, x, image.height), fill=(red, green, blue))
    draw.rectangle((44, 40, 436, 240), outline="#F2C14E", width=12)
    draw.ellipse((178, 78, 302, 202), fill="#F4F7FB", outline="#102A43", width=8)
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=False)
    stream.seek(0)
    return stream


def _insert_before_ext_lst(parent, child) -> None:
    ext_lst = parent.find(qn("a:extLst"))
    if ext_lst is None:
        parent.append(child)
    else:
        parent.insert(parent.index(ext_lst), child)


def _apply_bounded_shape3d(
    shape,
    *,
    light_rig: str = "threePt",
    light_rotation: tuple[int, int, int] | None = None,
    bevel_width_emu: int | None = 127000,
    bevel_height_emu: int | None = 127000,
    bevel_preset: str | None = "circle",
    explicit_zero_extrusion: bool = True,
    contour_width_emu: int | None = None,
    contour_color: str | None = None,
) -> None:
    """Insert the exact static 3D tuple used by the bounded renderer cohort."""
    sp_pr = shape._element.spPr
    scene3d = etree.Element(qn("a:scene3d"))
    etree.SubElement(scene3d, qn("a:camera"), prst="orthographicFront")
    light_rig_node = etree.SubElement(scene3d, qn("a:lightRig"), rig=light_rig, dir="t")
    if light_rotation is not None:
        lat, lon, rev = light_rotation
        etree.SubElement(
            light_rig_node,
            qn("a:rot"),
            lat=str(lat),
            lon=str(lon),
            rev=str(rev),
        )

    sp3d_attrs = {"extrusionH": "0"} if explicit_zero_extrusion else {}
    if contour_width_emu is not None:
        sp3d_attrs["contourW"] = str(contour_width_emu)
    sp3d = etree.Element(qn("a:sp3d"), **sp3d_attrs)
    bevel_attrs = {}
    if bevel_width_emu is not None:
        bevel_attrs["w"] = str(bevel_width_emu)
    if bevel_height_emu is not None:
        bevel_attrs["h"] = str(bevel_height_emu)
    if bevel_preset is not None:
        bevel_attrs["prst"] = bevel_preset
    etree.SubElement(sp3d, qn("a:bevelT"), **bevel_attrs)
    if contour_color is not None:
        contour_clr = etree.SubElement(sp3d, qn("a:contourClr"))
        etree.SubElement(contour_clr, qn("a:srgbClr"), val=contour_color)

    _insert_before_ext_lst(sp_pr, scene3d)
    _insert_before_ext_lst(sp_pr, sp3d)


def _apply_flat_shape3d_scene(
    shape,
    *,
    camera_preset: str,
    camera_rotation: tuple[int, int, int] | None = None,
    field_of_view: int | None = None,
    include_shape_format: bool = True,
) -> None:
    """Insert a camera-only scene around one zero-depth shape face."""
    sp_pr = shape._element.spPr
    camera_attrs = {"prst": camera_preset}
    if field_of_view is not None:
        camera_attrs["fov"] = str(field_of_view)
    scene3d = etree.Element(qn("a:scene3d"))
    camera = etree.SubElement(scene3d, qn("a:camera"), **camera_attrs)
    if camera_rotation is not None:
        lat, lon, rev = camera_rotation
        etree.SubElement(
            camera,
            qn("a:rot"),
            lat=str(lat),
            lon=str(lon),
            rev=str(rev),
        )
    etree.SubElement(scene3d, qn("a:lightRig"), rig="threePt", dir="t")
    _insert_before_ext_lst(sp_pr, scene3d)
    if include_shape_format:
        sp3d = etree.Element(qn("a:sp3d"), extrusionH="0")
        _insert_before_ext_lst(sp_pr, sp3d)


def _add_cjk_textbox(
    prs: Presentation,
    *,
    width: float = 8.5,
    height: float = 3.0,
    left: float = 1.0,
    top: float = 1.0,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(_emu(left), _emu(top), _emu(width), _emu(height))
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.12)
    text_frame.margin_right = Inches(0.12)
    text_frame.margin_top = Inches(0.08)
    text_frame.margin_bottom = Inches(0.08)
    return slide, shape, text_frame


def _rels_path(part_name: str) -> str:
    directory, filename = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _resolve_part_target(part_name: str, target: str) -> str:
    return posixpath.normpath(posixpath.join(posixpath.dirname(part_name), target))


def _relationship_target(zf: ZipFile, part_name: str, rel_suffix: str) -> str | None:
    root = etree.fromstring(zf.read(_rels_path(part_name)))
    for rel in root.xpath(".//pr:Relationship", namespaces=NS):
        rel_type = rel.get("Type", "")
        target = rel.get("Target")
        if target and rel_type == f"{PML_REL_PREFIX}{rel_suffix}":
            return _resolve_part_target(part_name, target)
    return None


def _read_pptx_entries(pptx_path: Path) -> tuple[list[tuple], dict[str, bytes]]:
    with ZipFile(pptx_path, "r") as zf:
        entries = [(info, zf.read(info.filename)) for info in zf.infolist()]
    return entries, {info.filename: data for info, data in entries}


def _replace_pptx_entries(pptx_path: Path, entries: list[tuple], patched: dict[str, bytes]) -> None:
    tmp_path = pptx_path.with_name(f"{pptx_path.name}.tmp")
    if tmp_path.exists():
        tmp_path.unlink()

    with ZipFile(tmp_path, "w", ZIP_DEFLATED) as out:
        for info, data in entries:
            out.writestr(info, patched.get(info.filename, data))
    tmp_path.replace(pptx_path)


def _patch_placeholder_idx_inheritance_case(pptx_path: Path) -> None:
    """Make slide placeholders inherit type through idx, matching real OOXML edge cases."""
    slide_part = "ppt/slides/slide1.xml"

    with ZipFile(pptx_path, "r") as zf:
        entries, data_by_name = _read_pptx_entries(pptx_path)
        layout_part = _relationship_target(zf, slide_part, "/slideLayout")
        if layout_part is None:
            raise RuntimeError("placeholder inheritance case slide has no slideLayout relationship")

        slide_root = etree.fromstring(data_by_name[slide_part])
        layout_root = etree.fromstring(data_by_name[layout_part])

    slide_title_ph = slide_root.xpath(".//p:ph[@type='title']", namespaces=NS)
    if not slide_title_ph:
        raise RuntimeError("placeholder inheritance case has no slide title placeholder")
    slide_title_ph[0].attrib.pop("type", None)
    slide_title_ph[0].set("idx", "0")

    slide_body_ph = slide_root.xpath(".//p:ph[@idx='1']", namespaces=NS)
    if not slide_body_ph:
        raise RuntimeError("placeholder inheritance case has no slide body placeholder")
    slide_body_ph[0].attrib.pop("type", None)

    layout_title_ph = layout_root.xpath(".//p:ph[@type='title']", namespaces=NS)
    if not layout_title_ph:
        raise RuntimeError("placeholder inheritance case has no layout title placeholder")
    layout_title_ph[0].set("idx", "0")

    patched = {
        slide_part: etree.tostring(slide_root, encoding="UTF-8", xml_declaration=True),
        layout_part: etree.tostring(layout_root, encoding="UTF-8", xml_declaration=True),
    }
    _replace_pptx_entries(pptx_path, entries, patched)


def _patch_connector_tail_arrows(pptx_path: Path) -> None:
    """Add OOXML tail arrowheads to generated connectors where python-pptx has no API."""
    slide_part = "ppt/slides/slide1.xml"
    entries, data_by_name = _read_pptx_entries(pptx_path)
    slide_root = etree.fromstring(data_by_name[slide_part])
    ns = {"p": PML_NS, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

    for ln in slide_root.xpath(".//p:cxnSp/p:spPr/a:ln", namespaces=ns):
        if ln.find(qn("a:tailEnd")) is not None:
            continue
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")

    _replace_pptx_entries(
        pptx_path,
        entries,
        {slide_part: etree.tostring(slide_root, encoding="UTF-8", xml_declaration=True)},
    )


def _patch_layered_transparency_case(pptx_path: Path) -> None:
    """Apply alpha modifiers to named overlap shapes for translucent-layer regressions."""
    slide_part = "ppt/slides/slide1.xml"
    entries, data_by_name = _read_pptx_entries(pptx_path)
    slide_root = etree.fromstring(data_by_name[slide_part])
    ns = {"p": PML_NS, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

    for sp in slide_root.xpath(".//p:sp", namespaces=ns):
        name = (sp.xpath("string(p:nvSpPr/p:cNvPr/@name)", namespaces=ns) or "").lower()
        if "alpha layer" not in name:
            continue
        srgb = sp.find(".//a:solidFill/a:srgbClr", namespaces=ns)
        if srgb is None or srgb.find(qn("a:alpha")) is not None:
            continue
        alpha = etree.SubElement(srgb, qn("a:alpha"))
        alpha.set("val", "45000")

    _replace_pptx_entries(
        pptx_path,
        entries,
        {slide_part: etree.tostring(slide_root, encoding="UTF-8", xml_declaration=True)},
    )


def _patch_scaled_group_diagram_case(pptx_path: Path) -> None:
    """Append a valid p:grpSp with non-identity chExt/ext scaling."""
    slide_part = "ppt/slides/slide1.xml"
    entries, data_by_name = _read_pptx_entries(pptx_path)
    slide_root = etree.fromstring(data_by_name[slide_part])
    sp_tree = slide_root.find(".//{%s}spTree" % PML_NS)
    if sp_tree is None:
        raise RuntimeError("scaled group case slide has no spTree")

    group_xml = f"""
      <p:grpSp xmlns:p="{PML_NS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <p:nvGrpSpPr>
          <p:cNvPr id="501" name="Scaled pipeline group"/>
          <p:cNvGrpSpPr/>
          <p:nvPr/>
        </p:nvGrpSpPr>
        <p:grpSpPr>
          <a:xfrm>
            <a:off x="914400" y="914400"/>
            <a:ext cx="9144000" cy="3657600"/>
            <a:chOff x="0" y="0"/>
            <a:chExt cx="4572000" cy="1828800"/>
          </a:xfrm>
        </p:grpSpPr>
        <p:sp>
          <p:nvSpPr><p:cNvPr id="502" name="Group Parse"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
          <p:spPr>
            <a:xfrm><a:off x="0" y="228600"/><a:ext cx="1219200" cy="685800"/></a:xfrm>
            <a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>
            <a:solidFill><a:srgbClr val="2F5597"/></a:solidFill>
          </p:spPr>
          <p:txBody>
            <a:bodyPr anchor="ctr"/><a:lstStyle/>
            <a:p><a:pPr algn="ctr"/><a:r><a:rPr sz="1800"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Parse</a:t></a:r></a:p>
          </p:txBody>
        </p:sp>
        <p:cxnSp>
          <p:nvCxnSpPr><p:cNvPr id="503" name="Group Arrow 1"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>
          <p:spPr>
            <a:xfrm><a:off x="1371600" y="571500"/><a:ext cx="457200" cy="0"/></a:xfrm>
            <a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>
            <a:ln w="25400"><a:solidFill><a:srgbClr val="5B9BD5"/></a:solidFill><a:tailEnd type="triangle" w="med" len="med"/></a:ln>
          </p:spPr>
        </p:cxnSp>
        <p:sp>
          <p:nvSpPr><p:cNvPr id="504" name="Group Model"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
          <p:spPr>
            <a:xfrm><a:off x="1828800" y="0"/><a:ext cx="1219200" cy="1143000"/></a:xfrm>
            <a:prstGeom prst="hexagon"><a:avLst/></a:prstGeom>
            <a:solidFill><a:srgbClr val="70AD47"/></a:solidFill>
          </p:spPr>
          <p:txBody>
            <a:bodyPr anchor="ctr"/><a:lstStyle/>
            <a:p><a:pPr algn="ctr"/><a:r><a:rPr sz="1800"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Model</a:t></a:r></a:p>
          </p:txBody>
        </p:sp>
        <p:cxnSp>
          <p:nvCxnSpPr><p:cNvPr id="505" name="Group Arrow 2"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>
          <p:spPr>
            <a:xfrm><a:off x="3200400" y="571500"/><a:ext cx="457200" cy="0"/></a:xfrm>
            <a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>
            <a:ln w="25400"><a:solidFill><a:srgbClr val="5B9BD5"/></a:solidFill><a:tailEnd type="triangle" w="med" len="med"/></a:ln>
          </p:spPr>
        </p:cxnSp>
        <p:sp>
          <p:nvSpPr><p:cNvPr id="506" name="Group Render"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
          <p:spPr>
            <a:xfrm><a:off x="3657600" y="228600"/><a:ext cx="914400" cy="685800"/></a:xfrm>
            <a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>
            <a:solidFill><a:srgbClr val="ED7D31"/></a:solidFill>
          </p:spPr>
          <p:txBody>
            <a:bodyPr anchor="ctr"/><a:lstStyle/>
            <a:p><a:pPr algn="ctr"/><a:r><a:rPr sz="1800"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Render</a:t></a:r></a:p>
          </p:txBody>
        </p:sp>
      </p:grpSp>
    """
    sp_tree.append(etree.fromstring(group_xml))

    _replace_pptx_entries(
        pptx_path,
        entries,
        {slide_part: etree.tostring(slide_root, encoding="UTF-8", xml_declaration=True)},
    )


# ---------------------------------------------------------------------------
# P0: Rich text cases
# ---------------------------------------------------------------------------

def _build_text_cases() -> list[CaseDef]:
    cases: list[CaseDef] = []
    seq = 0

    def _add(slug: str, build_fn, postprocess_fn=None, coverage=None):
        nonlocal seq
        seq += 1
        case = {
            "name": f"oracle-pypptx-text-{seq:04d}-{slug}",
            "build_fn": build_fn,
        }
        if postprocess_fn is not None:
            case["postprocess_fn"] = postprocess_fn
        if coverage is not None:
            case["coverage"] = coverage
        cases.append(case)

    # --- Font families ---
    font_families = [
        ("Arial", "arial"),
        ("Times New Roman", "times-new-roman"),
        ("Calibri", "calibri"),
        ("Courier New", "courier-new"),
        ("Georgia", "georgia"),
        ("Verdana", "verdana"),
        ("Impact", "impact"),
        ("Comic Sans MS", "comic-sans"),
    ]
    for font_name, slug in font_families:
        def _build(prs, _fn=font_name):
            sld = prs.slides.add_slide(prs.slide_layouts[6])  # blank
            txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(8), _emu(2))
            tf = txbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"The quick brown fox jumps over the lazy dog — {_fn}"
            run.font.name = _fn
            run.font.size = Pt(28)
        _add(f"font-{slug}", _build)

    # --- Font sizes ---
    for pt_size in [10, 14, 18, 24, 36, 48, 72]:
        def _build(prs, _sz=pt_size):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(10), _emu(3))
            tf = txbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"Font size {_sz}pt sample text"
            run.font.size = Pt(_sz)
            run.font.name = "Calibri"
        _add(f"size-{pt_size}pt", _build)

    # --- Bold / Italic / Underline combos ---
    style_combos = [
        ("bold", True, False, False),
        ("italic", False, True, False),
        ("underline", False, False, True),
        ("bold-italic", True, True, False),
        ("bold-underline", True, False, True),
        ("bold-italic-underline", True, True, True),
    ]
    for slug, bold, italic, underline in style_combos:
        def _build(prs, _b=bold, _i=italic, _u=underline, _s=slug):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(8), _emu(2))
            tf = txbox.text_frame
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"Style: {_s} — The quick brown fox"
            run.font.name = "Calibri"
            run.font.size = Pt(24)
            run.font.bold = _b
            run.font.italic = _i
            run.font.underline = _u
        _add(f"style-{slug}", _build)

    # --- Alignment ---
    alignments = [
        ("left", PP_ALIGN.LEFT),
        ("center", PP_ALIGN.CENTER),
        ("right", PP_ALIGN.RIGHT),
        ("justify", PP_ALIGN.JUSTIFY),
    ]
    for slug, align in alignments:
        def _build(prs, _a=align, _s=slug):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(8), _emu(4))
            tf = txbox.text_frame
            tf.word_wrap = True
            for i in range(3):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"Paragraph {i+1} aligned {_s}. Lorem ipsum dolor sit amet."
                p.alignment = _a
                p.font.size = Pt(18)
                p.font.name = "Calibri"
        _add(f"align-{slug}", _build)

    # --- Font colors ---
    colors = [
        ("red", RGBColor(0xFF, 0x00, 0x00)),
        ("green", RGBColor(0x00, 0xB0, 0x50)),
        ("blue", RGBColor(0x00, 0x70, 0xC0)),
        ("orange", RGBColor(0xFF, 0xC0, 0x00)),
        ("purple", RGBColor(0x7B, 0x2D, 0x8E)),
    ]
    for slug, color in colors:
        def _build(prs, _c=color, _s=slug):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(8), _emu(2))
            tf = txbox.text_frame
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"Color: {_s} text sample"
            run.font.name = "Calibri"
            run.font.size = Pt(28)
            run.font.color.rgb = _c
        _add(f"color-{slug}", _build)

    # --- Multi-paragraph with mixed formatting ---
    def _build_mixed(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(1), _emu(0.5), _emu(10), _emu(6))
        tf = txbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.LEFT
        r1 = p1.add_run()
        r1.text = "Title in Bold 36pt"
        r1.font.name = "Arial"
        r1.font.size = Pt(36)
        r1.font.bold = True

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = "Subtitle in italic 24pt — "
        r2.font.name = "Georgia"
        r2.font.size = Pt(24)
        r2.font.italic = True
        r3 = p2.add_run()
        r3.text = "with colored segment"
        r3.font.name = "Georgia"
        r3.font.size = Pt(24)
        r3.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)

        p3 = tf.add_paragraph()
        r4 = p3.add_run()
        r4.text = "Body text in Calibri 18pt. Lorem ipsum dolor sit amet, consectetur adipiscing elit."
        r4.font.name = "Calibri"
        r4.font.size = Pt(18)
    _add("mixed-formatting", _build_mixed)

    # --- Bullet list ---
    def _build_bullets(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(1), _emu(1), _emu(8), _emu(5))
        tf = txbox.text_frame
        tf.word_wrap = True
        items = [
            (0, "First level item A"),
            (1, "Second level item A.1"),
            (1, "Second level item A.2"),
            (0, "First level item B"),
            (1, "Second level item B.1"),
            (2, "Third level item B.1.a"),
            (0, "First level item C"),
        ]
        for i, (level, text) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.level = level
            p.font.size = Pt(18)
            p.font.name = "Calibri"
    _add("bullet-list", _build_bullets)

    # --- Vertical text (East Asian) ---
    def _build_vertical(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(4), _emu(0.5), _emu(3), _emu(6))
        tf = txbox.text_frame
        tf.word_wrap = True
        txBody = tf._txBody
        bodyPr = txBody.find(qn("a:bodyPr"))
        bodyPr.set("vert", "eaVert")
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "垂直文本テスト Vertical Text 수직 텍스트"
        run.font.size = Pt(24)
        run.font.name = "Microsoft YaHei"
    _add("vertical-east-asian", _build_vertical)

    # --- Vertical text (wordArtVert) ---
    def _build_vertical_word(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(4), _emu(0.5), _emu(2), _emu(6))
        tf = txbox.text_frame
        tf.word_wrap = True
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        bodyPr.set("vert", "wordArtVert")
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "VERTICAL STACKED TEXT"
        run.font.size = Pt(24)
        run.font.name = "Arial"
    _add("vertical-stacked", _build_vertical_word)

    # --- Text anchor (top / middle / bottom) ---
    _ANCHOR_XML = {"top": "t", "middle": "ctr", "bottom": "b"}
    for anchor_slug in ["top", "middle", "bottom"]:
        def _build(prs, _s=anchor_slug, _xml=_ANCHOR_XML[anchor_slug]):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            shp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(2), _emu(1), _emu(6), _emu(4))
            tf = shp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"Anchor: {_s}"
            run.font.size = Pt(24)
            run.font.name = "Calibri"
            run.font.bold = True
            p.alignment = PP_ALIGN.CENTER
            bodyPr = tf._txBody.find(qn("a:bodyPr"))
            bodyPr.set("anchor", _xml)
        _add(f"anchor-{anchor_slug}", _build)

    # --- Line spacing ---
    def _build_spacing(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(1), _emu(0.5), _emu(10), _emu(6))
        tf = txbox.text_frame
        tf.word_wrap = True
        for i in range(5):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"Line {i+1}: Lorem ipsum dolor sit amet, consectetur adipiscing elit."
            p.font.size = Pt(18)
            p.font.name = "Calibri"
            p.space_after = Pt(12)
            p.space_before = Pt(6)
    _add("line-spacing", _build_spacing)

    # --- Placeholder idx-only inheritance (slide ph idx -> layout/master ph type) ---
    def _build_placeholder_idx_inheritance(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content

        sld.shapes.title.text = "Placeholder Inheritance Title"

        body = sld.placeholders[1]
        tf = body.text_frame
        tf.clear()
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "Body placeholder inherits bullet and master text size"

        p2 = tf.add_paragraph()
        p2.text = "Second line also inherits body placeholder style"
        p2.level = 0
    _add(
        "placeholder-idx-inheritance",
        _build_placeholder_idx_inheritance,
        _patch_placeholder_idx_inheritance_case,
    )

    # --- CJK layout interaction matrix (issue #23 follow-up) ---
    # Keep these as focused one-feature variants. Their PPTX/PDF outputs stay local and
    # ignored; the definitions make the matrix reproducible across native PowerPoint hosts.
    cjk_wrap_text = "坚守问题导向，持续提升复杂演示文稿的渲染质量与一致性。"

    def _cjk_coverage(*features: str) -> dict:
        return {
            "oracle": "native-powerpoint",
            "requiredFonts": ["Microsoft YaHei"],
            "features": ["text.cjk", *features],
        }

    def _build_cjk_wrap_square_no_autofit(prs):
        _, _, tf = _add_cjk_textbox(prs, width=5.6, height=1.8)
        _configure_text_body(tf, wrap="square", autofit="noAutofit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-wrap-square-no-autofit",
        _build_cjk_wrap_square_no_autofit,
        coverage=_cjk_coverage("bodyPr.wrap=square", "bodyPr.noAutofit"),
    )

    def _build_cjk_wrap_square_implicit(prs):
        _, _, tf = _add_cjk_textbox(prs, width=5.6, height=1.8)
        _configure_text_body(tf, wrap="square", autofit=None)
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-wrap-square-implicit-autofit",
        _build_cjk_wrap_square_implicit,
        coverage=_cjk_coverage("bodyPr.wrap=square", "bodyPr.autofit=omitted"),
    )

    def _build_cjk_wrap_none_no_autofit(prs):
        _, _, tf = _add_cjk_textbox(prs, width=5.6, height=1.4)
        _configure_text_body(tf, wrap="none", autofit="noAutofit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-wrap-none-no-autofit",
        _build_cjk_wrap_none_no_autofit,
        coverage=_cjk_coverage("bodyPr.wrap=none", "bodyPr.noAutofit"),
    )

    def _build_cjk_sp_autofit(prs):
        _, _, tf = _add_cjk_textbox(prs, width=4.8, height=1.2)
        _configure_text_body(tf, wrap="square", autofit="spAutoFit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-sp-autofit-narrow",
        _build_cjk_sp_autofit,
        coverage=_cjk_coverage("bodyPr.wrap=square", "bodyPr.spAutoFit", "layout.narrow"),
    )

    def _build_cjk_norm_autofit(prs):
        _, _, tf = _add_cjk_textbox(prs, width=4.8, height=1.35)
        _configure_text_body(
            tf,
            wrap="square",
            autofit="normAutofit",
            autofit_attrs={"fontScale": "85000", "lnSpcReduction": "10000"},
        )
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-norm-autofit-scaled",
        _build_cjk_norm_autofit,
        coverage=_cjk_coverage(
            "bodyPr.wrap=square",
            "bodyPr.normAutofit",
            "normAutofit.fontScale=85000",
            "normAutofit.lnSpcReduction=10000",
        ),
    )

    multiline_cjk = "第一行：统一字体测量\v第二行：核对中文行距\v第三行：观察基线位置"

    def _build_cjk_line_spacing(prs, value_name: str, value: int):
        _, _, tf = _add_cjk_textbox(prs, width=9.0, height=4.2)
        _configure_text_body(tf, wrap="square", autofit="noAutofit")
        paragraph = tf.paragraphs[0]
        _set_cjk_paragraph_text(paragraph, multiline_cjk, font_size_pt=28)
        _replace_spacing_value(paragraph, "lnSpc", value_name, value)

    _add(
        "cjk-line-spacing-100pct",
        lambda prs: _build_cjk_line_spacing(prs, "spcPct", 100000),
        coverage=_cjk_coverage("paragraph.manualBreaks", "lnSpc.spcPct=100000"),
    )
    _add(
        "cjk-line-spacing-130pct",
        lambda prs: _build_cjk_line_spacing(prs, "spcPct", 130000),
        coverage=_cjk_coverage("paragraph.manualBreaks", "lnSpc.spcPct=130000"),
    )
    _add(
        "cjk-line-spacing-28pt",
        lambda prs: _build_cjk_line_spacing(prs, "spcPts", 2800),
        coverage=_cjk_coverage("paragraph.manualBreaks", "lnSpc.spcPts=2800"),
    )

    def _build_cjk_paragraph_spacing(prs, value_name: str, before: int, after: int):
        _, _, tf = _add_cjk_textbox(prs, width=9.0, height=4.8)
        _configure_text_body(tf, wrap="square", autofit="noAutofit")
        texts = [
            "第一段：段前段后间距需要遵循演示文稿定义。",
            "第二段：浏览器默认外边距不能参与布局。",
            "第三段：末段外边距需要保持 PowerPoint 语义。",
        ]
        for index, text in enumerate(texts):
            paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            _add_cjk_run(paragraph, text, font_size_pt=26)
            _replace_spacing_value(paragraph, "spcBef", value_name, before)
            _replace_spacing_value(paragraph, "spcAft", value_name, after)

    _add(
        "cjk-paragraph-spacing-points",
        lambda prs: _build_cjk_paragraph_spacing(prs, "spcPts", 600, 1000),
        coverage=_cjk_coverage("paragraph.multiple", "spcBef.spcPts", "spcAft.spcPts"),
    )
    _add(
        "cjk-paragraph-spacing-percent",
        lambda prs: _build_cjk_paragraph_spacing(prs, "spcPct", 30000, 50000),
        coverage=_cjk_coverage("paragraph.multiple", "spcBef.spcPct", "spcAft.spcPct"),
    )

    def _build_cjk_mixed_run_spacing(prs):
        _, _, tf = _add_cjk_textbox(prs, width=9.2, height=2.2)
        _configure_text_body(tf, wrap="square", autofit="noAutofit")
        paragraph = tf.paragraphs[0]
        _add_cjk_run(paragraph, "字距放宽", font_size_pt=32, bold=True, spacing=180)
        _add_cjk_run(paragraph, "｜正常字距｜", font_size_pt=32)
        _add_cjk_run(paragraph, "字距收紧", font_size_pt=32, bold=True, spacing=-120)
    _add(
        "cjk-mixed-run-character-spacing",
        _build_cjk_mixed_run_spacing,
        coverage=_cjk_coverage("runs.adjacent", "run.spacing=positive-negative"),
    )

    def _build_cjk_rounded_shape(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(2.0),
            _emu(1.2),
            _emu(8.5),
            _emu(4.2),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xE9, 0xF2, 0xFF)
        shape.line.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        _configure_text_body(
            tf,
            wrap="square",
            autofit="noAutofit",
            anchor="ctr",
        )
        paragraph = tf.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        _set_cjk_paragraph_text(
            paragraph,
            "容器内第一行\v容器内第二行\v居中与行距共同生效",
            font_size_pt=28,
            bold=True,
        )
        _replace_spacing_value(paragraph, "lnSpc", "spcPct", 120000)
    _add(
        "cjk-rounded-shape-centered-spacing",
        _build_cjk_rounded_shape,
        coverage=_cjk_coverage(
            "container.roundedRect",
            "bodyPr.anchor=ctr",
            "paragraph.alignment=center",
            "paragraph.manualBreaks",
            "lnSpc.spcPct=120000",
        ),
    )

    def _build_cjk_sp_autofit_square(prs):
        _, _, tf = _add_cjk_textbox(prs, width=2.4, height=2.4)
        _configure_text_body(tf, wrap="square", autofit="spAutoFit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-sp-autofit-square-growth",
        _build_cjk_sp_autofit_square,
        coverage=_cjk_coverage(
            "bodyPr.wrap=square",
            "bodyPr.spAutoFit",
            "layout.aspectRatio=square",
            "shape.growth",
        ),
    )

    def _build_cjk_sp_autofit_tall(prs):
        _, _, tf = _add_cjk_textbox(prs, width=1.8, height=3.2, top=0.6)
        _configure_text_body(tf, wrap="square", autofit="spAutoFit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-sp-autofit-tall-growth",
        _build_cjk_sp_autofit_tall,
        coverage=_cjk_coverage(
            "bodyPr.wrap=square",
            "bodyPr.spAutoFit",
            "layout.aspectRatio=tall",
            "shape.growth",
        ),
    )

    def _build_cjk_sp_autofit_wide(prs):
        _, _, tf = _add_cjk_textbox(prs, width=8.4, height=0.45)
        _configure_text_body(tf, wrap="square", autofit="spAutoFit")
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=30)
    _add(
        "cjk-sp-autofit-wide-compact",
        _build_cjk_sp_autofit_wide,
        coverage=_cjk_coverage(
            "bodyPr.wrap=square",
            "bodyPr.spAutoFit",
            "layout.aspectRatio=wide",
            "layout.compact",
        ),
    )

    def _build_cjk_sp_autofit_explicit_overflow(prs):
        _, _, tf = _add_cjk_textbox(prs, width=3.0, height=1.0)
        _configure_text_body(
            tf,
            wrap="none",
            autofit="spAutoFit",
            horz_overflow="overflow",
            vert_overflow="overflow",
        )
        _add_cjk_run(tf.paragraphs[0], cjk_wrap_text, font_size_pt=24)
    _add(
        "cjk-sp-autofit-explicit-overflow",
        _build_cjk_sp_autofit_explicit_overflow,
        coverage=_cjk_coverage(
            "bodyPr.wrap=none",
            "bodyPr.spAutoFit",
            "bodyPr.horzOverflow=overflow",
            "bodyPr.vertOverflow=overflow",
            "autofit.inverse-opt-out",
        ),
    )

    # --- Paragraph default-run color precedence (issues #21 / #23 follow-up) ---
    def _color_precedence_coverage(*features: str) -> dict:
        return {
            "oracle": "native-powerpoint",
            "requiredFonts": ["Microsoft YaHei"],
            "features": ["text.color.precedence", *features],
        }

    def _build_color_precedence_case(
        prs,
        *,
        width: float,
        height: float,
        left: float,
        top: float,
        text: str,
        default_color: tuple[str, str] | None,
        run_color: tuple[str, str] | None = None,
    ):
        _, shape, tf = _add_cjk_textbox(
            prs,
            width=width,
            height=height,
            left=left,
            top=top,
        )
        _configure_text_body(tf, wrap="square", autofit="noAutofit", anchor="ctr")
        paragraph = tf.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = _add_cjk_run(paragraph, text, font_size_pt=28, bold=True)
        _set_shape_font_ref_color(shape, "accent1")
        if default_color is not None:
            _set_paragraph_default_color(paragraph, *default_color)
        if run_color is not None:
            _set_run_explicit_color(run, *run_color)

    _add(
        "defrpr-srgb-over-fontref-square",
        lambda prs: _build_color_precedence_case(
            prs,
            width=3.4,
            height=3.4,
            left=5.0,
            top=1.7,
            text="段落默认色优先 SRGB RED",
            default_color=("srgbClr", "C00000"),
        ),
        coverage=_color_precedence_coverage(
            "shape.style.fontRef=accent1",
            "paragraph.defRPr.solidFill.srgbClr=C00000",
            "layout.aspectRatio=square",
            "precedence.paragraph-defRPr-over-shape-fontRef",
        ),
    )
    _add(
        "defrpr-scheme-over-fontref-wide",
        lambda prs: _build_color_precedence_case(
            prs,
            width=8.8,
            height=1.8,
            left=2.2,
            top=2.8,
            text="段落 schemeClr accent2 应覆盖 shape fontRef accent1",
            default_color=("schemeClr", "accent2"),
        ),
        coverage=_color_precedence_coverage(
            "shape.style.fontRef=accent1",
            "paragraph.defRPr.solidFill.schemeClr=accent2",
            "layout.aspectRatio=wide",
            "precedence.paragraph-defRPr-over-shape-fontRef",
        ),
    )
    _add(
        "run-color-over-defrpr-tall",
        lambda prs: _build_color_precedence_case(
            prs,
            width=2.4,
            height=5.0,
            left=5.5,
            top=1.0,
            text="显式 RUN 紫色覆盖段落默认橙色",
            default_color=("schemeClr", "accent2"),
            run_color=("srgbClr", "7030A0"),
        ),
        coverage=_color_precedence_coverage(
            "shape.style.fontRef=accent1",
            "paragraph.defRPr.solidFill.schemeClr=accent2",
            "run.rPr.solidFill.srgbClr=7030A0",
            "layout.aspectRatio=tall",
            "precedence.run-rPr-over-paragraph-defRPr",
        ),
    )
    _add(
        "fontref-fallback-no-defrpr",
        lambda prs: _build_color_precedence_case(
            prs,
            width=6.5,
            height=2.2,
            left=3.4,
            top=2.5,
            text="无显式颜色时使用 shape fontRef accent1",
            default_color=None,
        ),
        coverage=_color_precedence_coverage(
            "shape.style.fontRef=accent1",
            "paragraph.defRPr.color=absent",
            "run.rPr.color=absent",
            "precedence.fontRef-fallback",
            "precedence.inverse-opt-out",
        ),
    )

    def _build_styled_soft_break_matrix(prs):
        _, _, text_frame = _add_cjk_textbox(
            prs,
            width=9.5,
            height=5.0,
            left=1.8,
            top=1.1,
        )
        _configure_text_body(text_frame, wrap="square", autofit="noAutofit", anchor="t")

        for index, paragraph in enumerate(
            (text_frame.paragraphs[0], text_frame.add_paragraph())
        ):
            paragraph._p.remove(paragraph._p.get_or_add_endParaRPr())
            p_pr = paragraph._p.get_or_add_pPr()
            line_spacing = etree.SubElement(p_pr, qn("a:lnSpc"))
            etree.SubElement(line_spacing, qn("a:spcPct"), val="100000")
            if index == 1:
                etree.SubElement(p_pr, qn("a:buChar"), char="•")

            line_break = etree.SubElement(paragraph._p, qn("a:br"))
            break_props = etree.SubElement(line_break, qn("a:rPr"), sz="3000", lang="en-US")
            if index == 1:
                _replace_text_solid_fill(break_props, "srgbClr", "C00000")
            etree.SubElement(break_props, qn("a:latin"), typeface="Arial")

            run = paragraph.add_run()
            run.text = "Visible 10pt text after a 30pt styled soft break"
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run_props = run._r.get_or_add_rPr()
            run_props.set("lang", "en-US")
            if index == 1:
                _replace_text_solid_fill(run_props, "srgbClr", "0070C0")

    _add(
        "styled-soft-break-matrix",
        _build_styled_soft_break_matrix,
        coverage={
            "oracle": "native-powerpoint",
            "features": [
                "text.soft-break.rPr",
                "text.soft-break.font-size",
                "text.soft-break.font-family",
                "text.bullet.color-from-visible-run",
                "paragraph.line-spacing.spcPct=100000",
            ],
        },
    )

    def _build_tab_stop_matrix(prs):
        probe_left = 1.0
        probe_top = 2.1
        probe_width = 11.2
        probe_height = 2.0
        configs = [
            {
                "title": "Leading left tab with paragraph margin",
                "parts": ["\t", "段首目标文本"],
                "paragraph_attrs": {"marL": "457200", "algn": "l"},
                "tabs": [(1371600, "l")],
                "extra_guides": [(1828800, "margin + stop")],
            },
            {
                "title": "Inline left tab in a separate run",
                "parts": ["前缀", "\t", "行中目标文本"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(2743200, "l")],
            },
            {
                "title": "Multiple left tabs",
                "parts": ["A", "\t", "B", "\t", "C"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(1828800, "l"), (3657600, "l")],
            },
            {
                "title": "Bullet paragraph followed by a left tab",
                "parts": ["\t", "项目符号后的目标文本"],
                "paragraph_attrs": {"marL": "731520", "indent": "-274320", "algn": "l"},
                "tabs": [(1371600, "l")],
                "bullet": True,
            },
            {
                "title": "Inline left tab inside one text run",
                "parts": ["同一文本运行前缀\t目标文本"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(2743200, "l")],
            },
            {
                "title": "Default one-inch tab interval without tabLst",
                "parts": ["前缀", "\t", "默认目标文本"],
                "paragraph_attrs": {"defTabSz": "914400", "algn": "l"},
                "tabs": [],
                "extra_guides": [(914400, "1in"), (1828800, "2in"), (2743200, "3in")],
            },
            {
                "title": "Center-aligned explicit tab",
                "parts": ["前缀", "\t", "CENTER"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(3657600, "ctr")],
            },
            {
                "title": "Right-aligned explicit tab",
                "parts": ["前缀", "\t", "RIGHT"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(3657600, "r")],
            },
            {
                "title": "Decimal-aligned explicit tab",
                "parts": ["数值", "\t", "123.45"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(3657600, "dec")],
            },
            {
                "title": "Right-to-left explicit tab opt-out",
                "parts": ["\t", "مرحبا بالعالم"],
                "paragraph_attrs": {"rtl": "1", "algn": "l"},
                "tabs": [(2743200, "l")],
                "font_name": "Arial",
            },
            {
                "title": "Vertical East Asian explicit left tab",
                "parts": ["\t", "竖排制表目标"],
                "paragraph_attrs": {"algn": "l"},
                "tabs": [(1828800, "l")],
                "vertical": True,
            },
        ]

        for config in configs:
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            title_box = slide.shapes.add_textbox(
                _emu(0.7), _emu(0.35), _emu(11.9), _emu(0.65)
            )
            title_paragraph = title_box.text_frame.paragraphs[0]
            _add_cjk_run(
                title_paragraph,
                config["title"],
                font_name="Arial",
                font_size_pt=24,
                bold=True,
            )

            probe = slide.shapes.add_textbox(
                _emu(probe_left),
                _emu(probe_top),
                _emu(probe_width),
                _emu(probe_height),
            )
            probe.name = "Tab probe"
            probe.line.color.rgb = RGBColor(0x80, 0x80, 0x80)
            probe.line.width = Pt(1)
            text_frame = probe.text_frame
            text_frame.clear()
            text_frame.word_wrap = False
            text_frame.margin_left = 0
            text_frame.margin_right = 0
            text_frame.margin_top = 0
            text_frame.margin_bottom = 0
            _configure_text_body(text_frame, wrap="none", autofit="noAutofit", anchor="ctr")
            if config.get("vertical"):
                text_frame._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")

            paragraph = text_frame.paragraphs[0]
            p_pr = paragraph._p.get_or_add_pPr()
            for attr_name, attr_value in config["paragraph_attrs"].items():
                p_pr.set(attr_name, attr_value)
            if config.get("bullet"):
                etree.SubElement(p_pr, qn("a:buChar"), char="•")
            if config["tabs"]:
                tab_list = etree.SubElement(p_pr, qn("a:tabLst"))
                for position, alignment in config["tabs"]:
                    etree.SubElement(
                        tab_list,
                        qn("a:tab"),
                        pos=str(position),
                        algn=alignment,
                    )

            font_name = config.get("font_name", "Microsoft YaHei")
            for part in config["parts"]:
                _add_cjk_run(
                    paragraph,
                    part,
                    font_name=font_name,
                    font_size_pt=28,
                )

            guides = [(position, alignment) for position, alignment in config["tabs"]]
            guides.extend(config.get("extra_guides", []))
            for guide_index, (position, label) in enumerate(guides):
                if config.get("vertical"):
                    y = probe_top + position / 914400
                    line = slide.shapes.add_connector(
                        MSO_CONNECTOR.STRAIGHT,
                        _emu(probe_left - 0.3),
                        _emu(y),
                        _emu(probe_left + probe_width + 0.3),
                        _emu(y),
                    )
                    label_left = probe_left + 0.03
                    label_top = y + 0.03
                else:
                    x = probe_left + position / 914400
                    line = slide.shapes.add_connector(
                        MSO_CONNECTOR.STRAIGHT,
                        _emu(x),
                        _emu(1.45),
                        _emu(x),
                        _emu(4.65),
                    )
                    label_left = x + 0.03
                    label_top = 1.45
                line.line.color.rgb = (
                    RGBColor(0x00, 0x80, 0x00)
                    if guide_index == 0
                    else RGBColor(0xC0, 0x00, 0x00)
                )
                line.line.width = Pt(1)
                guide_label = slide.shapes.add_textbox(
                    _emu(label_left), _emu(label_top), _emu(1.4), _emu(0.35)
                )
                guide_run = guide_label.text_frame.paragraphs[0].add_run()
                guide_run.text = f"{position / 914400:g}in {label}"
                guide_run.font.name = "Arial"
                guide_run.font.size = Pt(9)

    _add(
        "tab-stop-matrix",
        _build_tab_stop_matrix,
        coverage={
            "oracle": "native-powerpoint",
            "requiredFonts": ["Microsoft YaHei", "Arial"],
            "features": [
                "text.tab.explicit-left-leading",
                "text.tab.explicit-left-inline",
                "text.tab.explicit-left-multiple",
                "text.tab.bullet",
                "text.tab.mixed-run",
                "text.tab.default-size",
                "text.tab.alignment=center|right|decimal",
                "text.tab.rtl=observation-only",
                "text.tab.vertical-left",
            ],
        },
    )
    cases[-1]["slide_count"] = 11

    return cases


# ---------------------------------------------------------------------------
# P1: Shape adjustment variants
# ---------------------------------------------------------------------------

def _build_shape_adj_cases() -> list[CaseDef]:
    cases: list[CaseDef] = []
    seq = 0

    def _add(slug: str, build_fn):
        nonlocal seq
        seq += 1
        cases.append({
            "name": f"oracle-pypptx-shape-adj-{seq:04d}-{slug}",
            "build_fn": build_fn,
        })

    # Shapes with meaningful adjustments: (MSO_SHAPE, slug, [(adj_index, value), ...])
    adj_configs = [
        # roundRect: corner radius
        (MSO_SHAPE.ROUNDED_RECTANGLE, "round-rect-small-radius", [(0, 0.05)]),
        (MSO_SHAPE.ROUNDED_RECTANGLE, "round-rect-large-radius", [(0, 0.45)]),
        # chevron: point depth
        (MSO_SHAPE.CHEVRON, "chevron-shallow", [(0, 0.15)]),
        (MSO_SHAPE.CHEVRON, "chevron-deep", [(0, 0.45)]),
        # right arrow: head width and depth
        (MSO_SHAPE.RIGHT_ARROW, "arrow-thin", [(0, 0.2), (1, 0.3)]),
        (MSO_SHAPE.RIGHT_ARROW, "arrow-wide-head", [(0, 0.1), (1, 0.6)]),
        # star 5-point: inner radius
        (MSO_SHAPE.STAR_5_POINT, "star5-thin", [(0, 0.15)]),
        (MSO_SHAPE.STAR_5_POINT, "star5-fat", [(0, 0.45)]),
        # donut: ring thickness
        (MSO_SHAPE.DONUT, "donut-thin-ring", [(0, 0.1)]),
        (MSO_SHAPE.DONUT, "donut-thick-ring", [(0, 0.45)]),
        # cross: arm thickness
        (MSO_SHAPE.CROSS, "cross-thin", [(0, 0.15)]),
        (MSO_SHAPE.CROSS, "cross-thick", [(0, 0.45)]),
        # trapezoid
        (MSO_SHAPE.TRAPEZOID, "trapezoid-narrow-top", [(0, 0.15)]),
        (MSO_SHAPE.TRAPEZOID, "trapezoid-wide-top", [(0, 0.45)]),
        # block arc
        (MSO_SHAPE.BLOCK_ARC, "block-arc-narrow", [(0, 0.1)]),
        (MSO_SHAPE.BLOCK_ARC, "block-arc-wide", [(0, 0.4)]),
        # folded corner
        (MSO_SHAPE.FOLDED_CORNER, "folded-corner-small", [(0, 0.1)]),
        (MSO_SHAPE.FOLDED_CORNER, "folded-corner-large", [(0, 0.4)]),
        # bevel
        (MSO_SHAPE.BEVEL, "bevel-thin", [(0, 0.05)]),
        (MSO_SHAPE.BEVEL, "bevel-thick", [(0, 0.35)]),
        # isosceles triangle: peak offset
        (MSO_SHAPE.ISOSCELES_TRIANGLE, "triangle-left-peak", [(0, 0.1)]),
        (MSO_SHAPE.ISOSCELES_TRIANGLE, "triangle-right-peak", [(0, 0.9)]),
        # pentagon
        (MSO_SHAPE.PENTAGON, "pentagon-shallow", [(0, 0.15)]),
        (MSO_SHAPE.PENTAGON, "pentagon-deep", [(0, 0.45)]),
        # can (cylinder): top ellipse height
        (MSO_SHAPE.CAN, "can-flat-top", [(0, 0.1)]),
        (MSO_SHAPE.CAN, "can-tall-top", [(0, 0.4)]),
        # heart
        (MSO_SHAPE.HEART, "heart-default", []),
        # moon
        (MSO_SHAPE.MOON, "moon-thin-crescent", [(0, 0.15)]),
        (MSO_SHAPE.MOON, "moon-wide-crescent", [(0, 0.7)]),
        # left brace
        (MSO_SHAPE.LEFT_BRACE, "left-brace-sharp", [(0, 0.05)]),
        (MSO_SHAPE.LEFT_BRACE, "left-brace-round", [(0, 0.3)]),
    ]

    for shape_type, slug, adjs in adj_configs:
        def _build(prs, _st=shape_type, _adjs=adjs):
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            shp = sld.shapes.add_shape(_st, _emu(2), _emu(1), _emu(5), _emu(4))
            for idx, val in _adjs:
                try:
                    shp.adjustments[idx] = val
                except (IndexError, ValueError) as exc:
                    print(f"    WARN: adj[{idx}]={val} failed on {_st}: {exc}", flush=True)
        _add(slug, _build)

    return cases


# ---------------------------------------------------------------------------
# P1a: Zero-adjustment flowchart geometry
# ---------------------------------------------------------------------------

def _build_flowchart_zero_adjustment_cases() -> list[CaseDef]:
    """Build native-oracle coverage for the 28 generated flowchart presets."""
    cases: list[CaseDef] = []
    configs = [
        (61, MSO_SHAPE.FLOWCHART_PROCESS, "process", "flowChartProcess"),
        (
            62,
            MSO_SHAPE.FLOWCHART_ALTERNATE_PROCESS,
            "alternate-process",
            "flowChartAlternateProcess",
        ),
        (63, MSO_SHAPE.FLOWCHART_DECISION, "decision", "flowChartDecision"),
        (64, MSO_SHAPE.FLOWCHART_DATA, "input-output", "flowChartInputOutput"),
        (
            65,
            MSO_SHAPE.FLOWCHART_PREDEFINED_PROCESS,
            "predefined-process",
            "flowChartPredefinedProcess",
        ),
        (
            66,
            MSO_SHAPE.FLOWCHART_INTERNAL_STORAGE,
            "internal-storage",
            "flowChartInternalStorage",
        ),
        (67, MSO_SHAPE.FLOWCHART_DOCUMENT, "document", "flowChartDocument"),
        (
            68,
            MSO_SHAPE.FLOWCHART_MULTIDOCUMENT,
            "multidocument",
            "flowChartMultidocument",
        ),
        (69, MSO_SHAPE.FLOWCHART_TERMINATOR, "terminator", "flowChartTerminator"),
        (70, MSO_SHAPE.FLOWCHART_PREPARATION, "preparation", "flowChartPreparation"),
        (71, MSO_SHAPE.FLOWCHART_MANUAL_INPUT, "manual-input", "flowChartManualInput"),
        (
            72,
            MSO_SHAPE.FLOWCHART_MANUAL_OPERATION,
            "manual-operation",
            "flowChartManualOperation",
        ),
        (73, MSO_SHAPE.FLOWCHART_CONNECTOR, "connector", "flowChartConnector"),
        (
            74,
            MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR,
            "offpage-connector",
            "flowChartOffpageConnector",
        ),
        (75, MSO_SHAPE.FLOWCHART_CARD, "punched-card", "flowChartPunchedCard"),
        (76, MSO_SHAPE.FLOWCHART_PUNCHED_TAPE, "punched-tape", "flowChartPunchedTape"),
        (
            77,
            MSO_SHAPE.FLOWCHART_SUMMING_JUNCTION,
            "summing-junction",
            "flowChartSummingJunction",
        ),
        (78, MSO_SHAPE.FLOWCHART_OR, "or", "flowChartOr"),
        (79, MSO_SHAPE.FLOWCHART_COLLATE, "collate", "flowChartCollate"),
        (80, MSO_SHAPE.FLOWCHART_SORT, "sort", "flowChartSort"),
        (81, MSO_SHAPE.FLOWCHART_EXTRACT, "extract", "flowChartExtract"),
        (82, MSO_SHAPE.FLOWCHART_MERGE, "merge", "flowChartMerge"),
        (83, MSO_SHAPE.FLOWCHART_STORED_DATA, "online-storage", "flowChartOnlineStorage"),
        (84, MSO_SHAPE.FLOWCHART_DELAY, "delay", "flowChartDelay"),
        (
            85,
            MSO_SHAPE.FLOWCHART_SEQUENTIAL_ACCESS_STORAGE,
            "magnetic-tape",
            "flowChartMagneticTape",
        ),
        (
            86,
            MSO_SHAPE.FLOWCHART_MAGNETIC_DISK,
            "magnetic-disk",
            "flowChartMagneticDisk",
        ),
        (
            87,
            MSO_SHAPE.FLOWCHART_DIRECT_ACCESS_STORAGE,
            "magnetic-drum",
            "flowChartMagneticDrum",
        ),
        (88, MSO_SHAPE.FLOWCHART_DISPLAY, "display", "flowChartDisplay"),
    ]

    def _style_explicit(shape) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x5B, 0x9B, 0xD5)
        shape.line.color.rgb = RGBColor(0x20, 0x38, 0x64)
        shape.line.width = Pt(2)

    def _use_solid_theme_reference(shape) -> None:
        style = shape._element.find(qn("p:style"))
        if style is None:
            raise RuntimeError("flowchart shape has no p:style")
        fill_ref = style.find(qn("a:fillRef"))
        if fill_ref is None:
            raise RuntimeError("flowchart shape style has no a:fillRef")
        fill_ref.set("idx", "1")
        scheme_color = fill_ref.find(qn("a:schemeClr"))
        if scheme_color is None:
            raise RuntimeError("flowchart shape fillRef has no a:schemeClr")
        scheme_color.set("val", "accent1")

    for shape_id, shape_type, slug, preset in configs:
        def _build(prs, _shape_type=shape_type, _preset=preset):
            square_slide = prs.slides.add_slide(prs.slide_layouts[6])
            square = square_slide.shapes.add_shape(
                _shape_type,
                _emu(4.5665),
                _emu(1.65),
                _emu(4.2),
                _emu(4.2),
            )
            square.name = f"{_preset} square explicit"
            _style_explicit(square)

            wide_slide = prs.slides.add_slide(prs.slide_layouts[6])
            wide = wide_slide.shapes.add_shape(
                _shape_type,
                _emu(2.6665),
                _emu(2.15),
                _emu(8.0),
                _emu(3.2),
            )
            wide.name = f"{_preset} wide theme"
            _use_solid_theme_reference(wide)

            tall_slide = prs.slides.add_slide(prs.slide_layouts[6])
            group = tall_slide.shapes.add_group_shape()
            tall = group.shapes.add_shape(
                _shape_type,
                _emu(1.0),
                _emu(1.0),
                _emu(4.0),
                _emu(4.0),
            )
            tall.name = f"{_preset} grouped tall explicit"
            _style_explicit(tall)
            group.left = _emu(5.0665)
            group.top = _emu(0.95)
            group.width = _emu(3.2)
            group.height = _emu(5.6)

        cases.append(
            {
                "name": f"oracle-pypptx-flowchart-{shape_id:04d}-{slug}",
                "build_fn": _build,
                "slide_count": 3,
                "coverage": {
                    "oracle": "native-powerpoint",
                    "features": [
                        f"p:sp.prstGeom={preset}",
                        "a:avLst.adjustmentGuideCount=0",
                        "geometry.aspect=square|wide|tall",
                        "container=standalone|nonIdentityGroup",
                        "paint=explicitSolid|themeStyleReference",
                    ],
                },
            }
        )

    return cases


# ---------------------------------------------------------------------------
# P1b: Bounded ordinary-shape effects
# ---------------------------------------------------------------------------


def _build_shape_effect_cases() -> list[CaseDef]:
    """Build native-oracle matrices for effects applied directly to ordinary shapes."""
    cases: list[CaseDef] = []

    def _set_simple_gradient(shape) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x3A, 0x7B, 0xD5)
        sp_pr = shape._element.spPr
        solid_fill = sp_pr.find(qn("a:solidFill"))
        if solid_fill is None:
            raise RuntimeError("shadow probe shape has no solid fill to replace")
        index = sp_pr.index(solid_fill)
        sp_pr.remove(solid_fill)
        gradient = etree.Element(qn("a:gradFill"), rotWithShape="1")
        stops = etree.SubElement(gradient, qn("a:gsLst"))
        first = etree.SubElement(stops, qn("a:gs"), pos="0")
        etree.SubElement(first, qn("a:srgbClr"), val="3A7BD5")
        second = etree.SubElement(stops, qn("a:gs"), pos="100000")
        etree.SubElement(second, qn("a:srgbClr"), val="74C0FC")
        etree.SubElement(gradient, qn("a:lin"), ang="5400000", scaled="1")
        sp_pr.insert(index, gradient)

    def _style_shape(shape, *, gradient: bool = False) -> None:
        if gradient:
            _set_simple_gradient(shape)
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x3A, 0x7B, 0xD5)
        shape.line.fill.background()

    def _apply_outer_shadow(
        shape,
        *,
        attributes: dict[str, int | str],
        color_kind: str = "srgbClr",
        color_value: str = "000000",
        color_modifiers: tuple[tuple[str, int], ...] = (("alpha", 35000),),
    ) -> None:
        sp_pr = shape._element.spPr
        _remove_children(sp_pr, ("effectLst", "effectDag"))
        effect_list = etree.Element(qn("a:effectLst"))
        shadow = etree.SubElement(
            effect_list,
            qn("a:outerShdw"),
            **{name: str(value) for name, value in attributes.items()},
        )
        color = etree.SubElement(shadow, qn(f"a:{color_kind}"), val=color_value)
        for modifier, value in color_modifiers:
            etree.SubElement(color, qn(f"a:{modifier}"), val=str(value))
        _insert_before_ext_lst(sp_pr, effect_list)

    def _add_standalone_shape(
        prs,
        *,
        shape_type,
        width: float,
        height: float,
        name: str,
        shadow_attributes: dict[str, int | str] | None,
        gradient: bool = False,
        shadow_color_kind: str = "srgbClr",
        shadow_color_value: str = "000000",
        shadow_color_modifiers: tuple[tuple[str, int], ...] = (("alpha", 35000),),
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            shape_type,
            _emu((13.333 - width) / 2),
            _emu((7.5 - height) / 2),
            _emu(width),
            _emu(height),
        )
        shape.name = name
        _style_shape(shape, gradient=gradient)
        if shadow_attributes is not None:
            _apply_outer_shadow(
                shape,
                attributes=shadow_attributes,
                color_kind=shadow_color_kind,
                color_value=shadow_color_value,
                color_modifiers=shadow_color_modifiers,
            )

    def _build_outer_shadow_matrix(prs) -> None:
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=5.2,
            height=3.2,
            name="No outer shadow inverse control",
            shadow_attributes=None,
        )
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=5.2,
            height=3.2,
            name="Rectangle blur with standard offset scale and alignment defaults",
            shadow_attributes={"blurRad": 127000, "rotWithShape": 0},
        )
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
            width=8.0,
            height=3.2,
            name="Wide rounded rectangle common offset shadow",
            shadow_attributes={
                "blurRad": 50800,
                "dist": 38100,
                "dir": 5400000,
                "rotWithShape": 0,
            },
        )
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.OVAL,
            width=3.2,
            height=5.4,
            name="Tall gradient ellipse directional shadow",
            shadow_attributes={
                "blurRad": 101600,
                "dist": 76200,
                "dir": 2700000,
                "algn": "ctr",
                "rotWithShape": 0,
            },
            gradient=True,
        )
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=5.2,
            height=3.2,
            name="Rectangle 102 percent centered shadow scale",
            shadow_attributes={
                "blurRad": 115455,
                "dist": 46182,
                "sx": 102000,
                "sy": 102000,
                "algn": "ctr",
                "rotWithShape": 0,
            },
        )
        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=5.2,
            height=3.2,
            name="Rectangle 92 percent top right scaled shadow",
            shadow_attributes={
                "blurRad": 317500,
                "dist": 127000,
                "dir": 8100000,
                "sx": 92000,
                "sy": 92000,
                "algn": "tr",
                "rotWithShape": 0,
            },
        )

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        grouped_shape = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(1.0),
            _emu(1.0),
            _emu(5.2),
            _emu(3.0),
        )
        grouped_shape.name = "Grouped rounded rectangle outer shadow"
        _style_shape(grouped_shape)
        _apply_outer_shadow(
            grouped_shape,
            attributes={
                "blurRad": 76200,
                "dist": 50800,
                "dir": 2700000,
                "rotWithShape": 0,
            },
        )
        group.left = _emu(3.4)
        group.top = _emu(1.55)
        group.width = _emu(6.5)
        group.height = _emu(3.75)

        _add_standalone_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=5.2,
            height=3.2,
            name="Scheme color outer shadow modifiers",
            shadow_attributes={
                "blurRad": 101600,
                "dist": 50800,
                "dir": 5400000,
                "sx": 100000,
                "sy": 100000,
                "algn": "b",
                "rotWithShape": 0,
            },
            shadow_color_kind="schemeClr",
            shadow_color_value="accent2",
            shadow_color_modifiers=(("lumMod", 60000), ("lumOff", 10000), ("alpha", 35000)),
        )

    cases.append(
        {
            "name": "oracle-pypptx-shape-effect-0001-outer-shadow-matrix",
            "build_fn": _build_outer_shadow_matrix,
            "slide_count": 8,
            "coverage": {
                "oracle": "native-powerpoint",
                "features": [
                    "p:sp/p:spPr/a:effectLst/a:outerShdw",
                    "semantics=inverse|defaults|offset|direction|uniformScale|colorModifiers",
                    "geometry=rect|roundRect|ellipse",
                    "geometry.aspect=square|wide|tall",
                    "container=standalone|singleLevelUniformUnrotatedGroup",
                    "paint=explicitSolid|simpleGradient",
                    "shadow.color=srgbClr|schemeClr+lumMod+lumOff+alpha",
                    "shadow.skew=absent",
                    "shape.transform=rotation0|flipHFalse|flipVFalse",
                    "shape3d=absent",
                ],
            },
            "assertions": {
                "inverseSlideIndices": [0],
                "positiveShadowSlideIndices": [1, 2, 3, 4, 5, 6, 7],
            },
        }
    )

    def _apply_reflection(
        shape,
        *,
        attributes: dict[str, int | str],
    ) -> None:
        sp_pr = shape._element.spPr
        _remove_children(sp_pr, ("effectLst", "effectDag"))
        effect_list = etree.Element(qn("a:effectLst"))
        etree.SubElement(
            effect_list,
            qn("a:reflection"),
            **{name: str(value) for name, value in attributes.items()},
        )
        _insert_before_ext_lst(sp_pr, effect_list)

    common_reflection = {
        "blurRad": 6350,
        "stA": 52000,
        "endA": 300,
        "endPos": 35000,
        "dir": 5400000,
        "sy": -100000,
        "algn": "bl",
        "rotWithShape": 0,
    }

    def _add_reflection_shape(
        prs,
        *,
        shape_type,
        width: float,
        height: float,
        name: str,
        reflection_attributes: dict[str, int | str] | None,
        gradient: bool = False,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            shape_type,
            _emu((13.333 - width) / 2),
            _emu(max(0.55, (7.5 - height * 1.65) / 2)),
            _emu(width),
            _emu(height),
        )
        shape.name = name
        _style_shape(shape, gradient=gradient)
        if reflection_attributes is not None:
            _apply_reflection(shape, attributes=reflection_attributes)

    def _build_reflection_matrix(prs) -> None:
        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=4.2,
            height=2.4,
            name="No reflection inverse control",
            reflection_attributes=None,
        )
        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=4.2,
            height=2.4,
            name="Solid rectangle common reflection",
            reflection_attributes=common_reflection,
        )
        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
            width=7.2,
            height=2.4,
            name="Wide rounded rectangle common reflection",
            reflection_attributes=common_reflection,
        )
        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.OVAL,
            width=2.8,
            height=3.4,
            name="Tall gradient ellipse common reflection",
            reflection_attributes=common_reflection,
            gradient=True,
        )
        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.UP_ARROW,
            width=4.0,
            height=2.5,
            name="Gradient arrow broad blur reflection",
            reflection_attributes={
                "blurRad": 177800,
                "stA": 40000,
                "endPos": 28000,
                "dir": 5400000,
                "sy": -100000,
                "algn": "bl",
                "rotWithShape": 0,
            },
            gradient=True,
        )

        grouped_slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = grouped_slide.shapes.add_group_shape()
        grouped_shape = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(1.0),
            _emu(1.0),
            _emu(4.2),
            _emu(2.4),
        )
        grouped_shape.name = "Grouped reflected rounded rectangle"
        _style_shape(grouped_shape)
        _apply_reflection(grouped_shape, attributes=common_reflection)
        group.left = _emu(4.0415)
        group.top = _emu(0.85)
        group.width = _emu(5.25)
        group.height = _emu(3.0)

        _add_reflection_shape(
            prs,
            shape_type=MSO_SHAPE.RECTANGLE,
            width=4.2,
            height=2.4,
            name="Offset reflection real corpus neighbor",
            reflection_attributes={
                "blurRad": 6350,
                "stA": 52000,
                "endA": 300,
                "endPos": 40000,
                "dist": 38100,
                "dir": 5400000,
                "sy": -100000,
                "algn": "bl",
                "rotWithShape": 0,
            },
        )

    def _build_text_reflection(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text_shape = slide.shapes.add_textbox(
            _emu(3.6665),
            _emu(1.6),
            _emu(6.0),
            _emu(1.4),
        )
        text_shape.name = "Live text reflection"
        text_shape.fill.background()
        text_shape.line.fill.background()
        text_frame = text_shape.text_frame
        text_frame.clear()
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = "LIVE REFLECTION"
        run.font.name = "Arial"
        run.font.size = Pt(34)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x20, 0x38, 0x64)
        _apply_reflection(
            text_shape,
            attributes={
                "blurRad": 12700,
                "stA": 35000,
                "endPos": 72000,
                "dir": 5400000,
                "sy": -100000,
                "algn": "bl",
                "rotWithShape": 0,
            },
        )

    cases.append(
        {
            "name": "oracle-pypptx-shape-effect-0002-reflection-matrix",
            "build_fn": _build_reflection_matrix,
            "slide_count": 7,
            "coverage": {
                "oracle": "native-powerpoint",
                "features": [
                    "p:sp/p:spPr/a:effectLst/a:reflection",
                    "semantics=inverse|alphaFade|blur|verticalFlip|distance",
                    "geometry=rect|roundRect|ellipse|upArrow",
                    "geometry.aspect=square|wide|tall",
                    "container=standalone|singleLevelUniformUnrotatedGroup",
                    "paint=explicitSolid|simpleGradient",
                    "reflection.direction=90deg",
                    "reflection.scale=sxImplicit100pct|syNegative100pct",
                    "reflection.alignment=bottomLeft",
                    "shape.transform=rotation0|flipHFalse|flipVFalse",
                    "shape3d=absent",
                ],
            },
            "assertions": {
                "inverseSlideIndices": [0],
                "positiveReflectionSlideIndices": [1, 2, 3, 4, 5, 6],
            },
        }
    )
    cases.append(
        {
            "name": "oracle-pypptx-text-effect-0001-reflection",
            "build_fn": _build_text_reflection,
            "coverage": {
                "oracle": "native-powerpoint",
                "features": [
                    "p:sp/p:spPr/a:effectLst/a:reflection",
                    "shape.kind=textBox",
                    "paint=noFillLiveText",
                    "bodyPr.wrap=none",
                    "bodyPr.autofit=spAutoFit",
                    "reflection.direction=90deg",
                    "reflection.scale=sxImplicit100pct|syNegative100pct",
                    "reflection.alignment=bottomLeft",
                    "shape.transform=rotation0|flipHFalse|flipVFalse",
                ],
            },
            "assertions": {"positiveReflectionSlideIndices": [0]},
        }
    )
    return cases


# ---------------------------------------------------------------------------
# P1c: Bounded static DrawingML 3D
# ---------------------------------------------------------------------------


def _build_shape3d_cases() -> list[CaseDef]:
    """Build a narrow native-oracle matrix for orthographic circle top bevels."""
    cases: list[CaseDef] = []
    seq = 0

    def _add(
        slug: str,
        build_fn,
        *,
        features: list[str],
        slide_count: int = 1,
        assertions: dict | None = None,
        claim: str | None = None,
    ):
        nonlocal seq
        seq += 1
        case = {
            "name": f"oracle-pypptx-shape3d-{seq:04d}-{slug}",
            "build_fn": build_fn,
            "coverage": {
                "oracle": "native-powerpoint",
                "features": features,
            },
        }
        if claim is not None:
            case["coverage"]["claim"] = claim
        if slide_count != 1:
            case["slide_count"] = slide_count
        if assertions is not None:
            case["assertions"] = assertions
        cases.append(case)

    def _add_picture(prs, *, apply_3d: bool):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu(3.0),
            _emu(1.75),
            _emu(7.333),
            _emu(4.0),
        )
        picture.name = "Static 3D picture"
        if apply_3d:
            # `twoPt:t` matches the real p:pic slice observed in model-platform.
            _apply_bounded_shape3d(picture, light_rig="twoPt")

    _add(
        "flat-optout",
        lambda prs: _add_picture(prs, apply_3d=False),
        features=["p:pic", "a:scene3d=absent", "a:sp3d=absent"],
    )
    _add(
        "picture-rect-circle-bevel",
        lambda prs: _add_picture(prs, apply_3d=True),
        features=[
            "p:pic",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=twoPt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _add_shape(
        prs,
        shape_type,
        *,
        left: float,
        top: float,
        width: float,
        height: float,
        contour: bool = False,
    ):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            shape_type,
            _emu(left),
            _emu(top),
            _emu(width),
            _emu(height),
        )
        shape.name = "Static 3D shape"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        shape.line.fill.background()
        _apply_bounded_shape3d(
            shape,
            contour_width_emu=12700 if contour else None,
            contour_color="FFFFFF" if contour else None,
        )

    _add(
        "roundrect-bevel-contour",
        lambda prs: _add_shape(
            prs,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left=3.0,
            top=1.75,
            width=7.333,
            height=4.0,
            contour=True,
        ),
        features=[
            "p:sp.prstGeom=roundRect",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT=circle",
            "a:sp3d.contourW=12700",
            "a:sp3d.contourClr=FFFFFF",
        ],
    )
    _add(
        "wide-bevel",
        lambda prs: _add_shape(
            prs,
            MSO_SHAPE.RECTANGLE,
            left=1.35,
            top=2.25,
            width=10.6,
            height=2.2,
        ),
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=wide",
            "a:scene3d.camera=orthographicFront",
            "a:sp3d.bevelT=circle",
        ],
    )
    _add(
        "tall-bevel",
        lambda prs: _add_shape(
            prs,
            MSO_SHAPE.RECTANGLE,
            left=5.15,
            top=0.55,
            width=3.0,
            height=6.4,
        ),
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=tall",
            "a:scene3d.camera=orthographicFront",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_grouped(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            _emu(1.0),
            _emu(1.0),
            _emu(4.0),
            _emu(2.5),
        )
        shape.name = "Grouped static 3D shape"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
        shape.line.fill.background()
        _apply_bounded_shape3d(shape)
        group.left = _emu(3.0)
        group.top = _emu(1.5)
        group.width = _emu(7.2)
        group.height = _emu(4.5)

    _add(
        "grouped-bevel",
        _build_grouped,
        features=[
            "p:grpSp/p:sp",
            "group.nonIdentityScale",
            "a:scene3d.camera=orthographicFront",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_real_picture_slice(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu(3.0),
            _emu(1.75),
            _emu(7.333),
            _emu(4.0),
        )
        picture.name = "Real-corpus static 3D picture slice"
        sp_pr = picture._element.spPr
        solid_fill = etree.Element(qn("a:solidFill"))
        white_fill = etree.SubElement(solid_fill, qn("a:srgbClr"), val="FFFFFF")
        etree.SubElement(white_fill, qn("a:shade"), val="85000")
        line = etree.Element(qn("a:ln"), w="88900", cap="sq")
        line_fill = etree.SubElement(line, qn("a:solidFill"))
        etree.SubElement(line_fill, qn("a:srgbClr"), val="FFFFFF")
        etree.SubElement(line, qn("a:miter"), lim="800000")
        effect_list = etree.Element(qn("a:effectLst"))
        shadow = etree.SubElement(
            effect_list,
            qn("a:outerShdw"),
            blurRad="55000",
            dist="18000",
            dir="5400000",
            algn="tl",
            rotWithShape="0",
        )
        shadow_color = etree.SubElement(shadow, qn("a:srgbClr"), val="000000")
        etree.SubElement(shadow_color, qn("a:alpha"), val="40000")
        for element in (solid_fill, line, effect_list):
            _insert_before_ext_lst(sp_pr, element)
        _apply_bounded_shape3d(
            picture,
            light_rig="twoPt",
            light_rotation=(0, 0, 7200000),
            bevel_width_emu=25400,
            bevel_height_emu=19050,
            bevel_preset=None,
            explicit_zero_extrusion=False,
            contour_color="FFFFFF",
        )

    _add(
        "real-picture-bevel-slice",
        _build_real_picture_slice,
        features=[
            "p:pic",
            "realCorpus=model-platform",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=twoPt:t",
            "a:scene3d.lightRig.rot=0,0,7200000",
            "a:sp3d.extrusionH=implicit-zero",
            "a:sp3d.bevelT=default-circle",
            "a:sp3d.bevelT.size=25400x19050",
            "a:sp3d.contourClr=FFFFFF",
            "a:effectLst.outerShdw=55000,18000,5400000",
        ],
    )

    def _add_cropped_picture(prs, crop: dict[str, float]) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu(3.0),
            _emu(1.4),
            _emu(7.333),
            _emu(4.7),
        )
        picture.name = "Cropped static 3D picture"
        for edge, value in crop.items():
            setattr(picture, f"crop_{edge}", value)
        _apply_bounded_shape3d(picture, light_rig="twoPt")

    _add(
        "picture-horizontal-crop-bevel",
        lambda prs: _add_cropped_picture(prs, {"left": 0.22, "right": 0.08}),
        features=[
            "p:pic",
            "a:srcRect=22%,0%,8%,0%",
            "crop.axis=horizontal",
            "a:scene3d.lightRig=twoPt:t",
            "a:sp3d.bevelT=circle",
        ],
    )
    _add(
        "picture-vertical-crop-bevel",
        lambda prs: _add_cropped_picture(prs, {"top": 0.18, "bottom": 0.12}),
        features=[
            "p:pic",
            "a:srcRect=0%,18%,0%,12%",
            "crop.axis=vertical",
            "a:scene3d.lightRig=twoPt:t",
            "a:sp3d.bevelT=circle",
        ],
    )
    _add(
        "picture-asymmetric-crop-bevel",
        lambda prs: _add_cropped_picture(
            prs,
            {"left": 0.12, "top": 0.08, "right": 0.18, "bottom": 0.10},
        ),
        features=[
            "p:pic",
            "a:srcRect=12%,8%,18%,10%",
            "crop.axis=combined",
            "a:scene3d.lightRig=twoPt:t",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_ellipse_matrix(prs):
        square_slide = prs.slides.add_slide(prs.slide_layouts[6])
        square = square_slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            _emu(4.5665),
            _emu(1.65),
            _emu(4.2),
            _emu(4.2),
        )
        square.name = "Static 3D ellipse square explicit"
        square.fill.solid()
        square.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        square.line.fill.background()
        _apply_bounded_shape3d(square)

        wide_slide = prs.slides.add_slide(prs.slide_layouts[6])
        wide = wide_slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            _emu(2.6665),
            _emu(2.15),
            _emu(8.0),
            _emu(3.2),
        )
        wide.name = "Static 3D ellipse wide theme"
        style = wide._element.find(qn("p:style"))
        if style is None:
            raise RuntimeError("ellipse shape has no p:style")
        fill_ref = style.find(qn("a:fillRef"))
        if fill_ref is None:
            raise RuntimeError("ellipse shape style has no a:fillRef")
        fill_ref.set("idx", "1")
        scheme_color = fill_ref.find(qn("a:schemeClr"))
        if scheme_color is None:
            raise RuntimeError("ellipse shape fillRef has no a:schemeClr")
        scheme_color.set("val", "accent1")
        wide.line.fill.background()
        _apply_bounded_shape3d(wide)

        tall_slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = tall_slide.shapes.add_group_shape()
        tall = group.shapes.add_shape(
            MSO_SHAPE.OVAL,
            _emu(1.0),
            _emu(1.0),
            _emu(4.0),
            _emu(4.0),
        )
        tall.name = "Static 3D ellipse grouped tall explicit"
        tall.fill.solid()
        tall.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
        tall.line.fill.background()
        _apply_bounded_shape3d(tall)
        group.left = _emu(5.0665)
        group.top = _emu(0.95)
        group.width = _emu(3.2)
        group.height = _emu(5.6)

    _add(
        "ellipse-circle-bevel-matrix",
        _build_ellipse_matrix,
        slide_count=3,
        features=[
            "p:sp.prstGeom=ellipse",
            "geometry.aspect=square|wide|tall",
            "container=standalone|nonIdentityGroup",
            "paint=explicitSolid|themeStyleReference",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _style_donut(shape, adjustment: float | None = None) -> None:
        if adjustment is not None:
            shape.adjustments[0] = adjustment
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        shape.line.fill.background()
        _apply_bounded_shape3d(shape)

    def _build_donut_matrix(prs):
        for name, adjustment in (
            ("Static 3D donut default", None),
            ("Static 3D donut lower bound", 0.0),
            ("Static 3D donut upper bound", 0.5),
        ):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            donut = slide.shapes.add_shape(
                MSO_SHAPE.DONUT,
                _emu(4.5665),
                _emu(1.65),
                _emu(4.2),
                _emu(4.2),
            )
            donut.name = name
            _style_donut(donut, adjustment)

        wide_slide = prs.slides.add_slide(prs.slide_layouts[6])
        wide = wide_slide.shapes.add_shape(
            MSO_SHAPE.DONUT,
            _emu(2.6665),
            _emu(2.15),
            _emu(8.0),
            _emu(3.2),
        )
        wide.name = "Static 3D donut wide adjusted theme"
        wide.adjustments[0] = 0.32
        style = wide._element.find(qn("p:style"))
        if style is None:
            raise RuntimeError("donut shape has no p:style")
        fill_ref = style.find(qn("a:fillRef"))
        if fill_ref is None:
            raise RuntimeError("donut shape style has no a:fillRef")
        fill_ref.set("idx", "1")
        scheme_color = fill_ref.find(qn("a:schemeClr"))
        if scheme_color is None:
            raise RuntimeError("donut shape fillRef has no a:schemeClr")
        scheme_color.set("val", "accent1")
        wide.line.fill.background()
        _apply_bounded_shape3d(wide)

        tall_slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = tall_slide.shapes.add_group_shape()
        tall = group.shapes.add_shape(
            MSO_SHAPE.DONUT,
            _emu(1.0),
            _emu(1.0),
            _emu(4.0),
            _emu(4.0),
        )
        tall.name = "Static 3D donut grouped tall adjusted"
        _style_donut(tall, 0.10)
        tall.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
        group.left = _emu(5.0665)
        group.top = _emu(0.95)
        group.width = _emu(3.2)
        group.height = _emu(5.6)

    _add(
        "donut-circle-bevel-adjustment-matrix",
        _build_donut_matrix,
        slide_count=5,
        features=[
            "p:sp.prstGeom=donut",
            "geometry.adjustment=0|10000|default25000|32000|50000",
            "geometry.aspect=square|wide|tall",
            "container=standalone|nonIdentityGroup",
            "paint=explicitSolid|themeStyleReference",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _style_camera_probe(shape, *, theme_fill: bool = False) -> None:
        if theme_fill:
            style = shape._element.find(qn("p:style"))
            if style is None:
                raise RuntimeError("camera probe shape has no p:style")
            fill_ref = style.find(qn("a:fillRef"))
            if fill_ref is None:
                raise RuntimeError("camera probe shape style has no a:fillRef")
            fill_ref.set("idx", "1")
            scheme_color = fill_ref.find(qn("a:schemeClr"))
            if scheme_color is None:
                raise RuntimeError("camera probe shape fillRef has no a:schemeClr")
            scheme_color.set("val", "accent1")
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        shape.line.fill.background()

    def _add_camera_probe(
        prs,
        *,
        camera_preset: str,
        camera_rotation: tuple[int, int, int] | None,
        field_of_view: int | None = None,
        width: float = 4.2,
        height: float = 4.2,
        theme_fill: bool = False,
        grouped: bool = False,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if grouped:
            group = slide.shapes.add_group_shape()
            shape = group.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                _emu(1.0),
                _emu(1.0),
                _emu(4.0),
                _emu(4.0),
            )
            group.left = _emu(3.8665)
            group.top = _emu(2.15)
            group.width = _emu(5.6)
            group.height = _emu(3.2)
        else:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                _emu((13.333 - width) / 2),
                _emu((7.5 - height) / 2),
                _emu(width),
                _emu(height),
            )
        shape.name = "Flat shape 3D camera projection probe"
        _style_camera_probe(shape, theme_fill=theme_fill)
        _apply_flat_shape3d_scene(
            shape,
            camera_preset=camera_preset,
            camera_rotation=camera_rotation,
            field_of_view=field_of_view,
        )

    def _build_camera_projection_matrix(prs) -> None:
        _add_camera_probe(
            prs,
            camera_preset="orthographicFront",
            camera_rotation=None,
        )
        _add_camera_probe(
            prs,
            camera_preset="orthographicFront",
            camera_rotation=(1200000, 1800000, 0),
        )
        for width, height, theme_fill, grouped in (
            (4.2, 4.2, False, False),
            (8.0, 3.2, True, False),
            (3.2, 5.4, False, False),
            (4.0, 4.0, False, True),
        ):
            _add_camera_probe(
                prs,
                camera_preset="perspectiveRelaxedModerately",
                camera_rotation=(18590633, 0, 0),
                field_of_view=7200000,
                width=width,
                height=height,
                theme_fill=theme_fill,
                grouped=grouped,
            )

    _add(
        "camera-projection-matrix",
        _build_camera_projection_matrix,
        slide_count=6,
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=square|wide|tall",
            "container=standalone|nonIdentityGroup",
            "paint=explicitSolid|themeStyleReference",
            "a:scene3d.camera=orthographicFront|perspectiveRelaxedModerately",
            "a:scene3d.camera.rot=absent|20deg,30deg,0deg|18590633,0,0",
            "a:scene3d.camera.fov=absent|7200000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevel=absent",
        ],
    )

    def _add_scene_only_probe(
        prs,
        *,
        width: float,
        height: float,
        text_plane: bool,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        left = _emu((13.333 - width) / 2)
        top = _emu((7.5 - height) / 2)
        if text_plane:
            shape = slide.shapes.add_textbox(left, top, _emu(width), _emu(height))
            shape.name = "Scene-only editable text plane"
            text_frame = shape.text_frame
            text_frame.clear()
            text_frame.margin_left = 0
            text_frame.margin_right = 0
            text_frame.margin_top = 0
            text_frame.margin_bottom = 0
            _configure_text_body(
                text_frame,
                wrap="none",
                autofit="spAutoFit",
                anchor="ctr",
            )
            for line_index, text in enumerate(("ZERO DEPTH", "EDITABLE TEXT", "NATIVE CAMERA")):
                paragraph = (
                    text_frame.paragraphs[0]
                    if line_index == 0
                    else text_frame.add_paragraph()
                )
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = text
                run.font.name = "Arial"
                run.font.size = Pt(26)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x20, 0x38, 0x64)
            _apply_flat_shape3d_scene(
                shape,
                camera_preset="perspectiveContrastingRightFacing",
                camera_rotation=(0, 19532225, 0),
                field_of_view=5100000,
                include_shape_format=False,
            )
        else:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                left,
                top,
                _emu(width),
                _emu(height),
            )
            shape.name = "Scene-only solid plane"
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
            shape.line.fill.background()
            _apply_flat_shape3d_scene(
                shape,
                camera_preset="perspectiveRelaxedModerately",
                camera_rotation=(18590633, 0, 0),
                field_of_view=7200000,
                include_shape_format=False,
            )

    def _build_scene_only_plane_matrix(prs) -> None:
        aspect_matrix = ((4.2, 4.2), (8.0, 3.2), (3.2, 5.4))
        for width, height in aspect_matrix:
            _add_scene_only_probe(prs, width=width, height=height, text_plane=False)
        for width, height in aspect_matrix:
            _add_scene_only_probe(prs, width=width, height=height, text_plane=True)

    _add(
        "scene-only-plane-matrix",
        _build_scene_only_plane_matrix,
        slide_count=6,
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=square|wide|tall",
            "surface=solidPlane|noFillTextPlane",
            "text.bodyPr=wrap-none|anchor-ctr|spAutoFit",
            "a:scene3d.camera=perspectiveRelaxedModerately|perspectiveContrastingRightFacing",
            "a:scene3d.camera.rot=18590633,0,0|0,19532225,0",
            "a:scene3d.camera.fov=7200000|5100000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d=absent",
            "effects=absent",
        ],
    )

    def _add_perspective_left_text_probe(
        prs,
        *,
        width: float,
        height: float,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(
            _emu((13.333 - width) / 2),
            _emu((7.5 - height) / 2),
            _emu(width),
            _emu(height),
        )
        shape.name = "Perspective-left editable text plane"
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0
        _configure_text_body(
            text_frame,
            wrap="none",
            autofit="spAutoFit",
        )
        body_pr = text_frame._txBody.find(qn("a:bodyPr"))
        if body_pr is None:
            raise RuntimeError("perspective-left text plane has no a:bodyPr")
        body_pr.attrib.pop("anchor", None)

        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = "透视 LEFT 120"
        run.font.name = "Arial"
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x20, 0x38, 0x64)

        _apply_flat_shape3d_scene(
            shape,
            camera_preset="perspectiveLeft",
            camera_rotation=None,
            field_of_view=7200000,
            include_shape_format=False,
        )

    def _build_perspective_left_text_plane_matrix(prs) -> None:
        for width, height in ((4.2, 4.2), (8.0, 3.2), (3.2, 5.4)):
            _add_perspective_left_text_probe(prs, width=width, height=height)

    _add(
        "perspective-left-text-plane-matrix",
        _build_perspective_left_text_plane_matrix,
        slide_count=3,
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=square|wide|tall",
            "surface=noFillTextPlane",
            "text.content=CJK|latin|digits",
            "text.bodyPr=wrap-none|anchor-absent|spAutoFit",
            "a:scene3d.camera=perspectiveLeft",
            "a:scene3d.camera.rot=absent",
            "a:scene3d.camera.fov=7200000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d=absent",
            "effects=absent",
        ],
    )

    def _add_perspective_right_picture_probe(
        prs,
        *,
        width: float,
        height: float,
        crop: dict[str, float],
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu((13.333 - width) / 2),
            _emu((7.5 - height) / 2),
            _emu(width),
            _emu(height),
        )
        picture.name = "Perspective-right editable picture plane"
        for edge, value in crop.items():
            setattr(picture, f"crop_{edge}", value)
        stretch = picture._element.find(qn("p:blipFill")).find(qn("a:stretch"))
        if stretch is None:
            raise RuntimeError("perspective-right picture plane has no a:stretch")
        _remove_children(stretch, ("fillRect",))
        _apply_flat_shape3d_scene(
            picture,
            camera_preset="perspectiveRight",
            camera_rotation=None,
            field_of_view=5700000,
            include_shape_format=False,
        )

    def _build_perspective_right_picture_plane_matrix(prs) -> None:
        for width, height, crop in (
            (4.2, 4.2, {}),
            (8.0, 3.2, {"left": 0.22, "right": 0.08}),
            (3.2, 5.4, {"top": 0.18, "bottom": 0.12}),
            (
                5.6,
                2.0,
                {"left": 0.01705, "top": 0.0335, "right": 0.01323, "bottom": 0.0335},
            ),
        ):
            _add_perspective_right_picture_probe(
                prs,
                width=width,
                height=height,
                crop=crop,
            )

    _add(
        "perspective-right-picture-plane-matrix",
        _build_perspective_right_picture_plane_matrix,
        slide_count=4,
        features=[
            "p:pic.prstGeom=rect",
            "geometry.aspect=square|wide|tall",
            "surface=stretchPngPicture",
            "a:srcRect=absent|horizontal|vertical|real-asymmetric",
            "a:scene3d.camera=perspectiveRight",
            "a:scene3d.camera.rot=absent",
            "a:scene3d.camera.fov=5700000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d=absent",
            "effects=absent",
        ],
    )

    def _add_default_top_bevel_probe(
        prs,
        *,
        shape_type,
        width: float,
        height: float,
        bevel_width_emu: int | None,
        bevel_height_emu: int | None,
        bevel_preset: str | None,
        name: str,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            shape_type,
            _emu((13.333 - width) / 2),
            _emu((7.5 - height) / 2),
            _emu(width),
            _emu(height),
        )
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        shape.line.fill.background()
        _apply_bounded_shape3d(
            shape,
            bevel_width_emu=bevel_width_emu,
            bevel_height_emu=bevel_height_emu,
            bevel_preset=bevel_preset,
        )

    def _build_default_top_bevel_dimensions_matrix(prs) -> None:
        for shape_type, width, height, omitted, explicit in (
            (
                MSO_SHAPE.RECTANGLE,
                4.2,
                4.2,
                (None, None, None, "Rect implicit bevel defaults"),
                (76200, 76200, "circle", "Rect explicit bevel defaults"),
            ),
            (
                MSO_SHAPE.ROUNDED_RECTANGLE,
                8.0,
                3.2,
                (None, 76200, "circle", "RoundRect implicit bevel width"),
                (76200, 76200, "circle", "RoundRect explicit bevel width"),
            ),
            (
                MSO_SHAPE.OVAL,
                3.2,
                5.4,
                (76200, None, "circle", "Ellipse implicit bevel height"),
                (76200, 76200, "circle", "Ellipse explicit bevel height"),
            ),
        ):
            for bevel_width, bevel_height, bevel_preset, name in (omitted, explicit):
                _add_default_top_bevel_probe(
                    prs,
                    shape_type=shape_type,
                    width=width,
                    height=height,
                    bevel_width_emu=bevel_width,
                    bevel_height_emu=bevel_height,
                    bevel_preset=bevel_preset,
                    name=name,
                )

    _add(
        "default-top-bevel-dimensions-matrix",
        _build_default_top_bevel_dimensions_matrix,
        slide_count=6,
        features=[
            "p:sp.prstGeom=rect|roundRect|ellipse",
            "geometry.aspect=square|wide|tall",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT.prst=implicit-circle|explicit-circle",
            "a:sp3d.bevelT.w=implicit-76200|explicit-76200",
            "a:sp3d.bevelT.h=implicit-76200|explicit-76200",
        ],
        assertions={
            "equivalentSlidePairs": [[0, 1], [2, 3], [4, 5]],
        },
    )

    def _build_donut_shadow_interpolation_matrix(prs) -> None:
        for aspect_label, width, height in (
            ("0.75", 3.75, 5.0),
            ("1.25", 5.0, 4.0),
            ("2.0", 7.0, 3.5),
        ):
            for adjustment_label, adjustment in (
                ("10000", 0.10),
                ("default25000", None),
                ("40000", 0.40),
            ):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                donut = slide.shapes.add_shape(
                    MSO_SHAPE.DONUT,
                    _emu((13.333 - width) / 2),
                    _emu((7.5 - height) / 2),
                    _emu(width),
                    _emu(height),
                )
                donut.name = (
                    f"Static 3D donut aspect {aspect_label} adjustment {adjustment_label}"
                )
                _style_donut(donut, adjustment)

    _add(
        "donut-shadow-interpolation-matrix",
        _build_donut_shadow_interpolation_matrix,
        slide_count=9,
        features=[
            "p:sp.prstGeom=donut",
            "geometry.adjustment=10000|default25000|40000",
            "geometry.aspect=0.75|1.25|2.0",
            "matrix=crossProduct(3x3)",
            "container=standalone",
            "paint=explicitSolid",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d.extrusionH=0",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _replace_with_multi_contour_cubic_geometry(shape) -> None:
        cust_geom = shape._element.spPr.find(qn("a:custGeom"))
        if cust_geom is None:
            raise RuntimeError("custom camera probe has no a:custGeom")
        path_list = cust_geom.find(qn("a:pathLst"))
        if path_list is None:
            raise RuntimeError("custom camera probe has no a:pathLst")
        for child in list(path_list):
            path_list.remove(child)

        path = etree.SubElement(path_list, qn("a:path"), w="1000", h="1000")

        def move_to(x: int, y: int) -> None:
            command = etree.SubElement(path, qn("a:moveTo"))
            etree.SubElement(command, qn("a:pt"), x=str(x), y=str(y))

        def line_to(x: int, y: int) -> None:
            command = etree.SubElement(path, qn("a:lnTo"))
            etree.SubElement(command, qn("a:pt"), x=str(x), y=str(y))

        def cubic_to(*points: tuple[int, int]) -> None:
            command = etree.SubElement(path, qn("a:cubicBezTo"))
            for x, y in points:
                etree.SubElement(command, qn("a:pt"), x=str(x), y=str(y))

        move_to(0, 500)
        cubic_to((0, 120), (280, 0), (450, 160))
        cubic_to((560, 270), (480, 500), (300, 580))
        line_to(0, 720)
        etree.SubElement(path, qn("a:close"))

        move_to(560, 180)
        cubic_to((700, 20), (1000, 120), (950, 430))
        cubic_to((920, 680), (680, 900), (520, 720))
        cubic_to((410, 590), (450, 320), (560, 180))
        etree.SubElement(path, qn("a:close"))

    def _add_perspective_custom_geometry_probe(
        prs,
        *,
        width: float,
        height: float,
        fill: tuple[int, int, int],
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if fill == (255, 255, 255):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(0x20, 0x38, 0x64)
        builder = slide.shapes.build_freeform(
            0,
            0,
            scale=(Inches(width) / 1000, Inches(height) / 1000),
        )
        builder.add_line_segments([(1000, 0), (1000, 1000), (0, 1000)], close=True)
        shape = builder.convert_to_shape(
            Inches((13.333 - width) / 2),
            Inches((7.5 - height) / 2),
        )
        shape.name = "Perspective custom geometry camera plane"
        _replace_with_multi_contour_cubic_geometry(shape)
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill)
        shape.line.fill.background()
        _apply_flat_shape3d_scene(
            shape,
            camera_preset="perspectiveRelaxedModerately",
            camera_rotation=(18590633, 0, 0),
            field_of_view=7200000,
            include_shape_format=False,
        )

    def _build_perspective_custom_geometry_plane_matrix(prs) -> None:
        for fill in ((47, 117, 181), (255, 255, 255)):
            for width, height in ((4.2, 4.2), (8.0, 3.2), (3.2, 5.4)):
                _add_perspective_custom_geometry_probe(
                    prs,
                    width=width,
                    height=height,
                    fill=fill,
                )

    _add(
        "perspective-custom-geometry-plane-matrix",
        _build_perspective_custom_geometry_plane_matrix,
        slide_count=6,
        features=[
            "p:sp.custGeom=multiContourCubic",
            "geometry.aspect=square|wide|tall",
            "matrix=crossProduct(3x2)",
            "container=standalone",
            "paint=explicitSolidBlue|explicitSolidWhite",
            "a:scene3d.camera=perspectiveRelaxedModerately",
            "a:scene3d.camera.rot=18590633,0,0",
            "a:scene3d.camera.fov=7200000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d=absent",
            "effects=absent",
        ],
    )

    def _apply_flat_group3d_scene(group) -> None:
        group_properties = group._element.grpSpPr
        _remove_children(group_properties, ("scene3d", "sp3d"))
        scene3d = etree.Element(qn("a:scene3d"))
        camera = etree.SubElement(
            scene3d,
            qn("a:camera"),
            prst="perspectiveLeft",
            fov="5700000",
        )
        etree.SubElement(camera, qn("a:rot"), lat="0", lon="1500000", rev="0")
        etree.SubElement(scene3d, qn("a:lightRig"), rig="threePt", dir="t")
        _insert_before_ext_lst(group_properties, scene3d)

    def _add_two_picture_group(
        shape_collection,
        *,
        left: float,
        top: float,
        width: float,
        height: float,
        apply_scene: bool,
        source_crop: dict[str, float] | None = None,
    ):
        group = shape_collection.add_group_shape()
        first = group.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu(0),
            _emu(0),
            _emu(4.0),
            _emu(1.2),
        )
        first.name = "Group camera picture upper"
        second = group.shapes.add_picture(
            _shape3d_fixture_image(),
            _emu(0),
            _emu(1.2),
            _emu(4.0),
            _emu(1.2),
        )
        second.name = "Group camera picture lower"
        if source_crop is not None:
            for picture in (first, second):
                for edge, value in source_crop.items():
                    setattr(picture, f"crop_{edge}", value)
        group.left = _emu(left)
        group.top = _emu(top)
        group.width = _emu(width)
        group.height = _emu(height)
        if apply_scene:
            _apply_flat_group3d_scene(group)
        return group

    def _build_perspective_left_picture_group_matrix(prs) -> None:
        for width, height, nested in (
            (4.2, 4.2, False),
            (8.0, 3.2, False),
            (3.2, 5.4, False),
            (5.0, 2.0, True),
        ):
            for apply_scene in (True, False):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                if nested:
                    outer = slide.shapes.add_group_shape()
                    _add_two_picture_group(
                        outer.shapes,
                        left=0.4,
                        top=0.3,
                        width=width,
                        height=height,
                        apply_scene=apply_scene,
                        source_crop={
                            "left": 0.01075,
                            "top": 0.41240,
                            "right": 0.01135,
                            "bottom": 0.41024,
                        },
                    )
                    outer.left = _emu((13.333 - width) / 2)
                    outer.top = _emu((7.5 - height) / 2)
                else:
                    _add_two_picture_group(
                        slide.shapes,
                        left=(13.333 - width) / 2,
                        top=(7.5 - height) / 2,
                        width=width,
                        height=height,
                        apply_scene=apply_scene,
                    )

    _add(
        "perspective-left-picture-group-matrix",
        _build_perspective_left_picture_group_matrix,
        slide_count=8,
        claim="native-verified-bounded-picture-group-camera-plane",
        assertions={
            "inverseSlideIndices": [1, 3, 5, 7],
            "positiveGroupCameraSlideIndices": [0, 2, 4, 6],
        },
        features=[
            "p:grpSp/p:grpSpPr/a:scene3d",
            "children=twoDirectStretchPngPictures",
            "pictureSourceCrop=absent|realCorpusAsymmetric",
            "geometry.aspect=square|wide|tall",
            "semantics=positive|sceneAbsentInverse",
            "container=standalone|coordinateOnlyAncestor",
            "a:scene3d.camera=perspectiveLeft",
            "a:scene3d.camera.rot=0,1500000,0",
            "a:scene3d.camera.fov=5700000",
            "a:scene3d.lightRig=threePt:t",
            "a:sp3d=absent",
            "targetGroup.effects=absent",
        ],
    )

    return cases


def _build_local_shape3d_cases() -> list[CaseDef]:
    """Build opt-in discovery probes without expanding the supported 3D cohort."""
    cases: list[CaseDef] = []
    seq = 0

    def _add(
        slug: str,
        build_fn,
        *,
        features: list[str],
        slide_count: int = 1,
        assertions: dict | None = None,
    ) -> None:
        nonlocal seq
        seq += 1
        case = {
            "name": f"oracle-local-shape3d-{seq:04d}-{slug}",
            "build_fn": build_fn,
            "local_only": True,
            "coverage": {
                "oracle": "native-powerpoint",
                "cohort": "experimental-local",
                "claim": "discovery-only",
                "features": features,
            },
        }
        if slide_count != 1:
            case["slide_count"] = slide_count
        if assertions is not None:
            case["assertions"] = assertions
        cases.append(case)

    def _style_probe(shape, *, color: tuple[int, int, int] = (47, 117, 181)) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.fill.background()
        _apply_bounded_shape3d(shape)

    def _build_preset(
        prs,
        shape_type,
        *,
        adjustment: float | None = None,
        width: float = 5.0,
        height: float = 5.0,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            shape_type,
            _emu((13.333 - width) / 2),
            _emu((7.5 - height) / 2),
            _emu(width),
            _emu(height),
        )
        shape.name = "Local static 3D discovery probe"
        if adjustment is not None:
            shape.adjustments[0] = adjustment
        _style_probe(shape)

    _add(
        "ellipse-circle-bevel",
        lambda prs: _build_preset(prs, MSO_SHAPE.OVAL),
        features=["p:sp.prstGeom=ellipse", "geometry.curved", "a:sp3d.bevelT=circle"],
    )
    _add(
        "donut-adjusted-circle-bevel",
        lambda prs: _build_preset(prs, MSO_SHAPE.DONUT, adjustment=0.32),
        features=[
            "p:sp.prstGeom=donut",
            "geometry.hole",
            "geometry.adjustment=0.32",
            "a:sp3d.bevelT=circle",
        ],
    )
    _add(
        "star5-adjusted-circle-bevel",
        lambda prs: _build_preset(
            prs,
            MSO_SHAPE.STAR_5_POINT,
            adjustment=0.36,
            width=5.6,
        ),
        features=[
            "p:sp.prstGeom=star5",
            "geometry.concave",
            "geometry.adjustment=0.36",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_freeform(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder = slide.shapes.build_freeform(
            0,
            0,
            scale=(Inches(5.2) / 1000, Inches(4.8) / 1000),
        )
        builder.add_line_segments(
            [
                (1000, 0),
                (1000, 390),
                (660, 390),
                (660, 1000),
                (340, 670),
                (0, 1000),
            ],
            close=True,
        )
        shape = builder.convert_to_shape(Inches(4.05), Inches(1.35))
        shape.name = "Concave freeform static 3D discovery probe"
        _style_probe(shape, color=(112, 173, 71))

    _add(
        "freeform-concave-circle-bevel",
        _build_freeform,
        features=["p:sp.custGeom", "geometry.concave", "a:sp3d.bevelT=circle"],
    )

    def _build_rotated(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(3.6),
            _emu(1.75),
            _emu(6.1),
            _emu(4.0),
        )
        shape.name = "Rotated static 3D discovery probe"
        shape.rotation = 30
        _style_probe(shape, color=(237, 125, 49))

    _add(
        "rotated-roundrect-bevel",
        _build_rotated,
        features=[
            "p:sp.prstGeom=roundRect",
            "shape.rotation=30deg",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_nested_group(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        shape = inner.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(0.8),
            _emu(0.6),
            _emu(3.8),
            _emu(2.4),
        )
        shape.name = "Nested grouped static 3D discovery probe"
        _style_probe(shape, color=(112, 48, 160))
        inner.left = _emu(1.0)
        inner.top = _emu(0.8)
        inner.width = _emu(5.8)
        inner.height = _emu(3.1)
        outer.left = _emu(2.25)
        outer.top = _emu(1.35)
        outer.width = _emu(8.6)
        outer.height = _emu(4.8)

    _add(
        "nested-group-scaled-bevel",
        _build_nested_group,
        features=[
            "p:grpSp/p:grpSp/p:sp",
            "group.nested",
            "group.nonIdentityScale",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _build_glow(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            _emu(3.0),
            _emu(1.75),
            _emu(7.333),
            _emu(4.0),
        )
        shape.name = "Glow and static 3D discovery probe"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x2F, 0x75, 0xB5)
        shape.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shape.line.width = Pt(2)
        effect_list = etree.Element(qn("a:effectLst"))
        glow = etree.SubElement(effect_list, qn("a:glow"), rad="114300")
        glow_color = etree.SubElement(glow, qn("a:srgbClr"), val="00B0F0")
        etree.SubElement(glow_color, qn("a:alpha"), val="65000")
        _insert_before_ext_lst(shape._element.spPr, effect_list)
        _apply_bounded_shape3d(shape)

    _add(
        "glow-roundrect-bevel",
        _build_glow,
        features=[
            "p:sp.prstGeom=roundRect",
            "a:effectLst.glow=114300",
            "a:ln=2pt",
            "a:sp3d.bevelT=circle",
        ],
    )

    def _apply_bottom_bevel(
        shape,
        *,
        bevel_preset: str = "relaxedInset",
        bevel_width_emu: int | None = None,
        bevel_height_emu: int | None = None,
        material: str | None = "dkEdge",
        light_rotation: tuple[int, int, int] | None = (0, 0, 3000000),
    ) -> None:
        sp_pr = shape._element.spPr
        scene3d = etree.Element(qn("a:scene3d"))
        etree.SubElement(scene3d, qn("a:camera"), prst="orthographicFront")
        light_rig = etree.SubElement(scene3d, qn("a:lightRig"), rig="threePt", dir="t")
        if light_rotation is not None:
            latitude, longitude, revolution = light_rotation
            etree.SubElement(
                light_rig,
                qn("a:rot"),
                lat=str(latitude),
                lon=str(longitude),
                rev=str(revolution),
            )

        shape_attrs = {"prstMaterial": material} if material is not None else {}
        sp3d = etree.Element(qn("a:sp3d"), **shape_attrs)
        bevel_attrs = {"prst": bevel_preset}
        if bevel_width_emu is not None:
            bevel_attrs["w"] = str(bevel_width_emu)
        if bevel_height_emu is not None:
            bevel_attrs["h"] = str(bevel_height_emu)
        etree.SubElement(sp3d, qn("a:bevelB"), **bevel_attrs)
        _insert_before_ext_lst(sp_pr, scene3d)
        _insert_before_ext_lst(sp_pr, sp3d)

    def _add_bottom_bevel_probe(
        prs,
        *,
        name: str,
        width: float = 6.0,
        height: float = 3.0,
        apply_3d: bool = True,
        bevel_preset: str = "relaxedInset",
        bevel_width_emu: int | None = None,
        bevel_height_emu: int | None = None,
        material: str | None = "dkEdge",
        light_rotation: tuple[int, int, int] | None = (0, 0, 3000000),
        text: str | None = None,
        transparent_overlay: bool = False,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        left = (13.333 - width) / 2
        top = (7.5 - height) / 2
        if transparent_overlay:
            base = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                _emu(left),
                _emu(top),
                _emu(width),
                _emu(height),
            )
            base.name = f"{name} base"
            base.fill.solid()
            base.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
            base.line.fill.background()

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            _emu(left),
            _emu(top),
            _emu(width),
            _emu(height),
        )
        shape.name = name
        shape.fill.solid()
        fill_color = (0xBD, 0xC4, 0xF0) if transparent_overlay else (0x44, 0x72, 0xC4)
        shape.fill.fore_color.rgb = RGBColor(*fill_color)
        shape.line.fill.background()
        if transparent_overlay:
            color = shape._element.spPr.find("a:solidFill/a:srgbClr", namespaces=shape._element.nsmap)
            if color is None:
                raise RuntimeError("bottom-bevel overlay has no solid color")
            etree.SubElement(color, qn("a:alpha"), val="5000")
            _insert_before_ext_lst(shape._element.spPr, etree.Element(qn("a:effectLst")))
        if apply_3d:
            _apply_bottom_bevel(
                shape,
                bevel_preset=bevel_preset,
                bevel_width_emu=bevel_width_emu,
                bevel_height_emu=bevel_height_emu,
                material=material,
                light_rotation=light_rotation,
            )
        if text is not None:
            text_frame = shape.text_frame
            text_frame.clear()
            body_pr = text_frame._txBody.find(qn("a:bodyPr"))
            if body_pr is not None:
                body_pr.set("anchor", "ctr")
            paragraph = text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def _build_bottom_relaxed_inset_matrix(prs) -> None:
        _add_bottom_bevel_probe(prs, name="Flat opaque control", apply_3d=False)
        _add_bottom_bevel_probe(prs, name="Bottom relaxedInset implicit defaults")
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset explicit defaults",
            bevel_width_emu=76200,
            bevel_height_emu=76200,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset live text",
            text="底部斜面",
        )
        _add_bottom_bevel_probe(
            prs,
            name="Real transparent bottom-bevel overlay",
            text="运营管理",
            transparent_overlay=True,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Flat transparent overlay control",
            apply_3d=False,
            text="运营管理",
            transparent_overlay=True,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset without material",
            material=None,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset without light rotation",
            light_rotation=None,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom circle neighbor",
            bevel_preset="circle",
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset square",
            width=4.2,
            height=4.2,
        )
        _add_bottom_bevel_probe(
            prs,
            name="Bottom relaxedInset tall",
            width=3.2,
            height=5.2,
        )

    _add(
        "bottom-relaxed-inset-matrix",
        _build_bottom_relaxed_inset_matrix,
        slide_count=11,
        assertions={"equivalentSlidePairs": [[1, 2], [1, 7], [1, 8]]},
        features=[
            "p:sp.prstGeom=rect",
            "geometry.aspect=square|wide|tall",
            "paint=opaqueSolid|transparentOverlay",
            "text=absent|liveCjk",
            "semantics=positive|exact-transparent-flat-control",
            "a:scene3d.camera=orthographicFront",
            "a:scene3d.lightRig=threePt:t",
            "a:scene3d.lightRig.rot=implicit|0,0,3000000",
            "a:sp3d.prstMaterial=implicit|dkEdge",
            "a:sp3d.bevelB.prst=relaxedInset|circle",
            "a:sp3d.bevelB.w=implicit-76200|explicit-76200",
            "a:sp3d.bevelB.h=implicit-76200|explicit-76200",
        ],
    )

    return cases


# ---------------------------------------------------------------------------
# P2: Native table interaction matrix
# ---------------------------------------------------------------------------

def _build_table_cases() -> list[CaseDef]:
    """Build isolated native-oracle cases for common DrawingML table semantics."""
    cases: list[CaseDef] = []
    seq = 0

    def _add(
        slug: str,
        build_fn,
        *,
        features: list[str],
        required_fonts: list[str] | None = None,
    ) -> None:
        nonlocal seq
        seq += 1
        coverage = {
            "oracle": "native-powerpoint",
            "features": features,
        }
        if required_fonts:
            coverage["requiredFonts"] = required_fonts
        cases.append(
            {
                "name": f"oracle-pypptx-table-{seq:04d}-{slug}",
                "build_fn": build_fn,
                "coverage": coverage,
            }
        )

    def _populate(table, values: list[list[str]], *, font_name: str = "Calibri") -> None:
        for row_idx, row in enumerate(values):
            for col_idx, value in enumerate(row):
                cell = table.cell(row_idx, col_idx)
                cell.text = value
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = font_name
                    paragraph.font.size = Pt(15 if row_idx == 0 else 13)
                    if row_idx == 0:
                        paragraph.font.bold = True

    def _build_default_grid(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(4, 4, _emu(1.2), _emu(1.0), _emu(10.8), _emu(4.8)).table
        _populate(
            table,
            [
                ["Quarter", "Revenue", "Cost", "Profit"],
                ["Q1", "$120K", "$85K", "$35K"],
                ["Q2", "$145K", "$92K", "$53K"],
                ["Q3", "$132K", "$88K", "$44K"],
            ],
        )

    _add(
        "default-grid",
        _build_default_grid,
        features=["a:tbl.defaultStyle", "a:tblGrid.equalColumns", "a:tr.equalRows"],
    )

    def _build_header_banded_rows(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(6, 4, _emu(0.9), _emu(0.7), _emu(11.4), _emu(5.6)).table
        table.first_row = True
        table.horz_banding = True
        _populate(
            table,
            [
                ["Region", "Target", "Actual", "Status"],
                ["North", "80", "83", "On track"],
                ["South", "75", "69", "Review"],
                ["East", "90", "94", "On track"],
                ["West", "70", "72", "On track"],
                ["Central", "65", "61", "Review"],
            ],
        )

    _add(
        "header-banded-rows",
        _build_header_banded_rows,
        features=["a:tblPr.firstRow=1", "a:tblPr.bandRow=1", "a:tableStyleId"],
    )

    def _build_first_last_columns(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(5, 5, _emu(0.7), _emu(1.0), _emu(11.9), _emu(4.8)).table
        table.first_row = True
        table.first_col = True
        table.last_col = True
        _populate(
            table,
            [
                ["Team", "Jan", "Feb", "Mar", "Total"],
                ["Alpha", "12", "14", "13", "39"],
                ["Beta", "9", "11", "15", "35"],
                ["Gamma", "16", "15", "17", "48"],
                ["Delta", "10", "12", "11", "33"],
            ],
        )

    _add(
        "first-last-columns",
        _build_first_last_columns,
        features=["a:tblPr.firstRow=1", "a:tblPr.firstCol=1", "a:tblPr.lastCol=1"],
    )

    def _build_horizontal_vertical_merges(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(5, 5, _emu(0.8), _emu(0.8), _emu(11.7), _emu(5.6)).table
        _populate(table, [[f"R{r + 1}C{c + 1}" for c in range(5)] for r in range(5)])
        table.cell(0, 0).merge(table.cell(0, 2))
        table.cell(0, 0).text = "Horizontal span ×3"
        table.cell(1, 3).merge(table.cell(3, 3))
        table.cell(1, 3).text = "Vertical\nspan ×3"
        table.cell(4, 0).merge(table.cell(4, 1))
        table.cell(4, 0).text = "Footer span"

    _add(
        "horizontal-vertical-merges",
        _build_horizontal_vertical_merges,
        features=["a:tc.gridSpan", "a:tc.rowSpan", "a:tc.hMerge", "a:tc.vMerge"],
    )

    def _build_variable_grid_sizes(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(3, 4, _emu(1.4), _emu(1.2), _emu(10.5), _emu(3.6)).table
        for column, width in zip(table.columns, (1.5, 3.0, 2.0, 4.0), strict=True):
            column.width = _emu(width)
        for row, height in zip(table.rows, (0.6, 1.2, 1.8), strict=True):
            row.height = _emu(height)
        _populate(
            table,
            [
                ["ID", "Variable width label", "State", "Long notes column"],
                ["A-01", "Standard row", "Ready", "The middle row is twice the header height."],
                ["A-02", "Tall row wraps across lines", "Review", "The final row is three times the header height."],
            ],
        )

    _add(
        "variable-grid-sizes",
        _build_variable_grid_sizes,
        features=["a:gridCol.variableWidth", "a:tr.variableHeight", "a:tc.textWrap"],
    )

    def _build_cell_margins_vertical_anchors(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(3, 3, _emu(1.3), _emu(0.8), _emu(10.7), _emu(5.8)).table
        _populate(
            table,
            [
                ["Top / 0.10 in", "Default peer", "Short"],
                ["Middle / 0.20 in", "Default peer", "Two\nlines"],
                ["Bottom / 0.30 in", "Default peer", "Three\ntext\nlines"],
            ],
        )
        for row_idx, (anchor, margin) in enumerate(
            (
                (MSO_VERTICAL_ANCHOR.TOP, 0.1),
                (MSO_VERTICAL_ANCHOR.MIDDLE, 0.2),
                (MSO_VERTICAL_ANCHOR.BOTTOM, 0.3),
            )
        ):
            cell = table.cell(row_idx, 0)
            cell.vertical_anchor = anchor
            cell.margin_left = Inches(margin)

    _add(
        "cell-margins-vertical-anchors",
        _build_cell_margins_vertical_anchors,
        features=["a:tcPr.anchor=t|ctr|b", "a:tcPr.marL=91440|182880|274320"],
    )

    def _build_cjk_mixed_text(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(4, 3, _emu(1.0), _emu(0.9), _emu(11.3), _emu(5.2)).table
        _populate(
            table,
            [
                ["项目", "状态 Status", "说明"],
                ["渲染精度", "进行中", "中文标点：括号（）、引号“”、顿号、。"],
                ["Chart 2D", "95.6%", "Latin + 中文 + 12345"],
                ["公式", "计划支持", "分数、根式、上下标与矩阵"],
            ],
            font_name="Microsoft YaHei",
        )

    _add(
        "cjk-mixed-text",
        _build_cjk_mixed_text,
        features=["a:tbl.cjkText", "a:tbl.mixedScript", "a:tbl.punctuationWrap"],
        required_fonts=["Microsoft YaHei"],
    )

    def _set_cell_border(
        cell,
        side: str,
        *,
        width: int,
        color: str | None = None,
        dash: str | None = None,
        no_fill: bool = False,
    ) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tag = qn(f"a:ln{side}")
        existing = tc_pr.find(tag)
        if existing is not None:
            tc_pr.remove(existing)
        line = etree.SubElement(tc_pr, tag, w=str(width))
        if no_fill:
            etree.SubElement(line, qn("a:noFill"))
        else:
            fill = etree.SubElement(line, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), val=color or "000000")
        if dash:
            etree.SubElement(line, qn("a:prstDash"), val=dash)

    def _build_cell_border_matrix(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(3, 3, _emu(1.4), _emu(1.0), _emu(10.4), _emu(4.8)).table
        _populate(
            table,
            [
                ["Red left 2 pt", "Blue dashed top 3 pt", "No right border"],
                ["Green bottom 1 pt", "Four explicit edges", "Theme peers"],
                ["Default", "Default", "Default"],
            ],
        )
        _set_cell_border(table.cell(0, 0), "L", width=25400, color="C00000")
        _set_cell_border(table.cell(0, 1), "T", width=38100, color="4472C4", dash="dash")
        _set_cell_border(table.cell(0, 2), "R", width=12700, no_fill=True)
        _set_cell_border(table.cell(1, 0), "B", width=12700, color="70AD47")
        for side in ("L", "R", "T", "B"):
            _set_cell_border(table.cell(1, 1), side, width=19050, color="7030A0")

    _add(
        "cell-border-matrix",
        _build_cell_border_matrix,
        features=["a:tcPr.lnL|lnR|lnT|lnB", "a:ln.solidFill", "a:ln.noFill", "a:ln.prstDash"],
    )

    return cases


# ---------------------------------------------------------------------------
# P2a: Native formula / OMML matrix
# ---------------------------------------------------------------------------

def _math_run(text: str, *, normal: bool = False) -> str:
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    math_properties = '<m:rPr><m:nor m:val="1"/></m:rPr>' if normal else ""
    return (
        f"<m:r>{math_properties}"
        '<a:rPr lang="en-US" sz="3200"><a:latin typeface="Cambria Math"/></a:rPr>'
        f"<m:t>{escaped}</m:t></m:r>"
    )


def _math_arg(name: str, content: str) -> str:
    return f"<m:{name}>{content}</m:{name}>"


def _patch_formula_case(pptx_path: Path, omml_content: str) -> None:
    """Wrap one fallback text box in the PowerPoint DrawingML math MCE shape."""
    slide_part = "ppt/slides/slide1.xml"
    entries, data_by_name = _read_pptx_entries(pptx_path)
    root = etree.fromstring(data_by_name[slide_part])
    fallback_shapes = root.xpath(
        ".//p:sp[p:nvSpPr/p:cNvPr[@name='Equation fallback']]",
        namespaces={"p": PML_NS},
    )
    if len(fallback_shapes) != 1:
        raise RuntimeError("formula case requires exactly one Equation fallback shape")

    fallback_shape = fallback_shapes[0]
    choice_shape = etree.fromstring(etree.tostring(fallback_shape))
    choice_name = choice_shape.xpath("./p:nvSpPr/p:cNvPr", namespaces={"p": PML_NS})[0]
    choice_name.set("name", "Native equation")
    paragraphs = choice_shape.xpath("./p:txBody/a:p", namespaces={"p": PML_NS, "a": DRAWINGML_NS})
    if len(paragraphs) != 1:
        raise RuntimeError("formula choice requires exactly one text paragraph")
    paragraph = paragraphs[0]
    for child in list(paragraph):
        if etree.QName(child).localname != "pPr":
            paragraph.remove(child)

    math_wrapper = etree.fromstring(
        (
            f'<a14:m xmlns:a14="{A14_NS}" xmlns:a="{DRAWINGML_NS}" '
            f'xmlns:m="{OMML_NS}"><m:oMathPara><m:oMath>{omml_content}'
            "</m:oMath></m:oMathPara></a14:m>"
        ).encode("utf-8")
    )
    paragraph.append(math_wrapper)
    etree.SubElement(paragraph, f"{{{DRAWINGML_NS}}}endParaRPr", lang="en-US")

    alternate = etree.Element(
        f"{{{MC_NS}}}AlternateContent",
        nsmap={"mc": MC_NS, "a14": A14_NS, "m": OMML_NS},
    )
    choice = etree.SubElement(alternate, f"{{{MC_NS}}}Choice", Requires="a14")
    choice.append(choice_shape)
    fallback = etree.SubElement(alternate, f"{{{MC_NS}}}Fallback")
    parent = fallback_shape.getparent()
    if parent is None:
        raise RuntimeError("formula fallback shape has no parent")
    parent.replace(fallback_shape, alternate)
    fallback.append(fallback_shape)

    _replace_pptx_entries(
        pptx_path,
        entries,
        {slide_part: etree.tostring(root, encoding="UTF-8", xml_declaration=True)},
    )


def _build_formula_cases() -> list[CaseDef]:
    """Build a bounded OMML construct matrix with visible fallback controls."""
    cases: list[CaseDef] = []

    def _add(slug: str, fallback_text: str, omml_content: str, feature: str) -> None:
        def _build(prs, _fallback_text=fallback_text):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(_emu(1.4), _emu(1.6), _emu(10.5), _emu(3.2))
            box.name = "Equation fallback"
            text_frame = box.text_frame
            text_frame.clear()
            text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            paragraph = text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = _fallback_text
            run.font.name = "Cambria Math"
            run.font.size = Pt(32)

        cases.append(
            {
                "name": f"oracle-pypptx-formula-{len(cases) + 1:04d}-{slug}",
                "build_fn": _build,
                "postprocess_fn": lambda path, _omml=omml_content: _patch_formula_case(
                    path, _omml
                ),
                "coverage": {
                    "oracle": "native-powerpoint",
                    "features": ["mc:AlternateContent", "a14:m", "m:oMathPara", feature],
                    "requiredFonts": ["Cambria Math"],
                },
            }
        )

    _add(
        "inline-expression",
        "x + 1",
        _math_run("x") + _math_run("+") + _math_run("1"),
        "m:r",
    )
    _add(
        "fraction",
        "(a + b) / (c + d)",
        "<m:f><m:fPr><m:type m:val=\"bar\"/></m:fPr>"
        + _math_arg("num", _math_run("a+b"))
        + _math_arg("den", _math_run("c+d"))
        + "</m:f>",
        "m:f",
    )
    _add(
        "radical",
        "sqrt(x^2 + y^2)",
        '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/>'
        + _math_arg("e", _math_run("x²+y²"))
        + "</m:rad>",
        "m:rad",
    )
    _add(
        "subscript-superscript",
        "x_i^2",
        "<m:sSubSup><m:sSubSupPr/>"
        + _math_arg("e", _math_run("x"))
        + _math_arg("sub", _math_run("i"))
        + _math_arg("sup", _math_run("2"))
        + "</m:sSubSup>",
        "m:sSubSup",
    )
    _add(
        "delimiters",
        "(x + y)",
        '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
        + _math_arg("e", _math_run("x+y"))
        + "</m:d>",
        "m:d",
    )
    _add(
        "nary-summation",
        "sum(i=1..n) i",
        '<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>'
        "</m:naryPr>"
        + _math_arg("sub", _math_run("i=1"))
        + _math_arg("sup", _math_run("n"))
        + _math_arg("e", _math_run("i"))
        + "</m:nary>",
        "m:nary",
    )
    _add(
        "matrix-2x2",
        "[1 2; 3 4]",
        "<m:m><m:mr>"
        + _math_arg("e", _math_run("1"))
        + _math_arg("e", _math_run("2"))
        + "</m:mr><m:mr>"
        + _math_arg("e", _math_run("3"))
        + _math_arg("e", _math_run("4"))
        + "</m:mr></m:m>",
        "m:m",
    )
    _add(
        "function",
        "sin(theta)",
        "<m:func><m:funcPr/>"
        + _math_arg("fName", _math_run("sin", normal=True))
        + _math_arg("e", _math_run("θ"))
        + "</m:func>",
        "m:func",
    )
    return cases


# ---------------------------------------------------------------------------
# P3: Composite (multi-component) cases
# ---------------------------------------------------------------------------

def _build_composite_cases() -> list[CaseDef]:
    cases: list[CaseDef] = []
    seq = 0

    def _add(slug: str, build_fn, postprocess_fn=None):
        nonlocal seq
        seq += 1
        case = {
            "name": f"oracle-pypptx-composite-{seq:04d}-{slug}",
            "build_fn": build_fn,
        }
        if postprocess_fn is not None:
            case["postprocess_fn"] = postprocess_fn
        cases.append(case)

    def _style_shape_text(shp, text: str, size: int = 16, color=RGBColor(0xFF, 0xFF, 0xFF)):
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = color

    # --- Two shapes side by side ---
    def _build_two_shapes(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(1), _emu(1.5), _emu(4), _emu(3))
        sld.shapes.add_shape(MSO_SHAPE.OVAL, _emu(6.5), _emu(1.5), _emu(4), _emu(3))
    _add("two-shapes-side-by-side", _build_two_shapes)

    # --- Shape with text inside ---
    def _build_shape_with_text(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        shp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(2), _emu(1), _emu(6), _emu(4))
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Text inside a rounded rectangle"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(24)
        p.font.name = "Calibri"
        p.font.bold = True
    _add("shape-with-centered-text", _build_shape_with_text)

    # --- Shape + Textbox overlay ---
    def _build_shape_textbox_overlay(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(1.5), _emu(1), _emu(7), _emu(5))
        txbox = sld.shapes.add_textbox(_emu(2), _emu(2), _emu(6), _emu(3))
        tf = txbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = "Overlaid Title"
        p1.alignment = PP_ALIGN.CENTER
        p1.font.size = Pt(32)
        p1.font.bold = True
        p2 = tf.add_paragraph()
        p2.text = "Body text overlaid on a rectangle shape"
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(18)
    _add("shape-textbox-overlay", _build_shape_textbox_overlay)

    # --- Multiple shapes (grid) ---
    def _build_shape_grid(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = [MSO_SHAPE.RECTANGLE, MSO_SHAPE.OVAL, MSO_SHAPE.DIAMOND,
                  MSO_SHAPE.HEXAGON, MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.HEART,
                  MSO_SHAPE.CROSS, MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.DONUT]
        for i, st in enumerate(shapes):
            row, col = divmod(i, 3)
            left = _emu(1 + col * 3.5)
            top = _emu(0.5 + row * 2.2)
            sld.shapes.add_shape(st, left, top, _emu(2.5), _emu(1.8))
    _add("shape-grid-3x3", _build_shape_grid)

    # --- Table + textbox ---
    def _build_table_textbox(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        # Title textbox
        txbox = sld.shapes.add_textbox(_emu(1), _emu(0.3), _emu(8), _emu(0.8))
        p = txbox.text_frame.paragraphs[0]
        p.text = "Quarterly Results"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.name = "Calibri"
        # Table
        tbl_shape = sld.shapes.add_table(4, 4, _emu(1), _emu(1.5), _emu(8), _emu(4))
        tbl = tbl_shape.table
        headers = ["Quarter", "Revenue", "Cost", "Profit"]
        data = [
            ["Q1", "$120K", "$85K", "$35K"],
            ["Q2", "$145K", "$92K", "$53K"],
            ["Q3", "$132K", "$88K", "$44K"],
        ]
        for j, h in enumerate(headers):
            cell = tbl.cell(0, j)
            cell.text = h
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(14)
        for i, row in enumerate(data):
            for j, val in enumerate(row):
                tbl.cell(i + 1, j).text = val
    _add("table-with-title", _build_table_textbox)

    # --- Chart + title textbox ---
    def _build_chart_title(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        txbox = sld.shapes.add_textbox(_emu(1), _emu(0.3), _emu(8), _emu(0.8))
        p = txbox.text_frame.paragraphs[0]
        p.text = "Sales Overview"
        p.font.size = Pt(28)
        p.font.bold = True
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
        chart_data.add_series("Product A", (45, 52, 48, 61))
        chart_data.add_series("Product B", (32, 38, 41, 35))
        sld.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            _emu(1), _emu(1.5), _emu(8), _emu(5),
            chart_data,
        )
    _add("chart-with-title", _build_chart_title)

    # --- Multiple textboxes with different styles ---
    def _build_multi_text(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        configs = [
            (_emu(0.5), _emu(0.5), _emu(4), _emu(2), "Arial", 20, True, PP_ALIGN.LEFT, "Left-aligned Arial Bold"),
            (_emu(5.5), _emu(0.5), _emu(5), _emu(2), "Georgia", 20, False, PP_ALIGN.RIGHT, "Right-aligned Georgia Italic"),
            (_emu(0.5), _emu(3), _emu(10), _emu(2), "Calibri", 16, False, PP_ALIGN.CENTER, "Centered Calibri — Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
            (_emu(0.5), _emu(5.5), _emu(10), _emu(1.5), "Courier New", 14, False, PP_ALIGN.LEFT, "Monospace: code_sample = function(x) { return x * 2; }"),
        ]
        for left, top, w, h, font, size, bold, align, text in configs:
            txbox = sld.shapes.add_textbox(left, top, w, h)
            tf = txbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.alignment = align
            p.font.name = font
            p.font.size = Pt(size)
            p.font.bold = bold
            if font == "Georgia":
                p.font.italic = True
    _add("multi-textbox-styles", _build_multi_text)

    # --- Two charts side by side ---
    def _build_two_charts(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd1 = CategoryChartData()
        cd1.categories = ["Jan", "Feb", "Mar"]
        cd1.add_series("Sales", (120, 135, 148))
        sld.shapes.add_chart(XL_CHART_TYPE.LINE, _emu(0.5), _emu(0.5), _emu(5.5), _emu(6), cd1)

        cd2 = CategoryChartData()
        cd2.categories = ["A", "B", "C", "D"]
        cd2.add_series("Share", (35, 25, 22, 18))
        sld.shapes.add_chart(XL_CHART_TYPE.PIE, _emu(6.5), _emu(0.5), _emu(5.5), _emu(6), cd2)
    _add("two-charts-line-pie", _build_two_charts)

    # --- Shapes with different fills + text ---
    def _build_colored_shapes_text(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        configs = [
            (MSO_SHAPE.ROUNDED_RECTANGLE, _emu(0.5), _emu(1), RGBColor(0x00, 0x70, 0xC0), "Blue Box"),
            (MSO_SHAPE.OVAL, _emu(4.5), _emu(1), RGBColor(0xFF, 0x40, 0x40), "Red Oval"),
            (MSO_SHAPE.HEXAGON, _emu(8.5), _emu(1), RGBColor(0x00, 0xB0, 0x50), "Green Hex"),
        ]
        for st, left, top, color, text in configs:
            shp = sld.shapes.add_shape(st, left, top, _emu(3.5), _emu(4))
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            tf = shp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = text
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _add("colored-shapes-with-text", _build_colored_shapes_text)

    # --- Table + chart + text ---
    def _build_dashboard(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        # Title
        txbox = sld.shapes.add_textbox(_emu(0.5), _emu(0.2), _emu(12), _emu(0.7))
        p = txbox.text_frame.paragraphs[0]
        p.text = "Dashboard: Monthly KPIs"
        p.font.size = Pt(28)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        # Small table (left)
        tbl_shape = sld.shapes.add_table(3, 2, _emu(0.5), _emu(1.2), _emu(4), _emu(2.5))
        tbl = tbl_shape.table
        for j, h in enumerate(["Metric", "Value"]):
            tbl.cell(0, j).text = h
        tbl.cell(1, 0).text = "Users"
        tbl.cell(1, 1).text = "12,450"
        tbl.cell(2, 0).text = "Revenue"
        tbl.cell(2, 1).text = "$89K"
        # Chart (right)
        cd = CategoryChartData()
        cd.categories = ["Mon", "Tue", "Wed", "Thu", "Fri"]
        cd.add_series("Visits", (850, 920, 780, 1050, 990))
        sld.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            _emu(5), _emu(1.2), _emu(7.5), _emu(5.5), cd,
        )
    _add("dashboard-table-chart", _build_dashboard)

    # --- Process flow with arrowed connectors and callouts ---
    def _build_process_flow_connectors(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        title = sld.shapes.add_textbox(_emu(0.5), _emu(0.25), _emu(12), _emu(0.5))
        p = title.text_frame.paragraphs[0]
        p.text = "Pipeline with connectors, labels, and callouts"
        p.font.size = Pt(24)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        steps = [
            ("Ingest", 0x2F5597),
            ("Parse", 0x70AD47),
            ("Model", 0xFFC000),
            ("Render", 0xED7D31),
        ]
        boxes = []
        for idx, (label, color) in enumerate(steps):
            shp = sld.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                _emu(0.8 + idx * 3.1),
                _emu(2.2),
                _emu(2.2),
                _emu(1.0),
            )
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)
            shp.line.color.rgb = RGBColor(0x2F, 0x2F, 0x2F)
            _style_shape_text(shp, label, 18)
            boxes.append(shp)

        for idx in range(len(boxes) - 1):
            x1 = _emu(3.0 + idx * 3.1)
            x2 = _emu(3.85 + idx * 3.1)
            conn = sld.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, _emu(2.7), x2, _emu(2.7))
            conn.line.width = Pt(2.25)
            conn.line.color.rgb = RGBColor(0x5B, 0x9B, 0xD5)

        callout = sld.shapes.add_shape(MSO_SHAPE.CLOUD_CALLOUT, _emu(4.3), _emu(4.2), _emu(4.5), _emu(1.3))
        callout.fill.solid()
        callout.fill.fore_color.rgb = RGBColor(0xFF, 0xF2, 0xCC)
        callout.line.color.rgb = RGBColor(0xBF, 0x90, 0x00)
        _style_shape_text(callout, "Connectors should stay stroke-only with arrowheads", 14, RGBColor(0, 0, 0))
    _add("process-flow-connectors", _build_process_flow_connectors, _patch_connector_tail_arrows)

    # --- Merged table with adjacent callouts and direct cell styling ---
    def _build_merged_table_callouts(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        tbl_shape = sld.shapes.add_table(5, 4, _emu(0.8), _emu(1.0), _emu(7.2), _emu(4.7))
        tbl = tbl_shape.table
        tbl.cell(0, 0).merge(tbl.cell(0, 3))
        tbl.cell(0, 0).text = "Merged table header"
        for p in tbl.cell(0, 0).text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(16)
            p.alignment = PP_ALIGN.CENTER
        for r in range(1, 5):
            tbl.cell(r, 0).text = f"Phase {r}"
            tbl.cell(r, 1).text = f"{70 + r * 4}%"
            tbl.cell(r, 2).text = "OK" if r % 2 else "Review"
            tbl.cell(r, 3).text = f"T+{r}"
        tbl.cell(2, 2).merge(tbl.cell(3, 2))
        tbl.cell(2, 2).text = "Merged\nstatus"

        for col, color in [(0, RGBColor(0xD9, 0xE2, 0xF3)), (3, RGBColor(0xE2, 0xF0, 0xD9))]:
            for r in range(1, 5):
                cell = tbl.cell(r, col)
                cell.fill.solid()
                cell.fill.fore_color.rgb = color

        note = sld.shapes.add_shape(MSO_SHAPE.LINE_CALLOUT_2, _emu(8.4), _emu(1.35), _emu(3.8), _emu(2.2))
        note.fill.solid()
        note.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
        note.line.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
        _style_shape_text(note, "Merged cells plus callout geometry", 14, RGBColor(0, 0, 0))
    _add("merged-table-callouts", _build_merged_table_callouts)

    # --- Rotated text and shapes mixed with curved connectors ---
    def _build_rotated_text_and_shapes(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        left = sld.shapes.add_shape(MSO_SHAPE.TRAPEZOID, _emu(1.0), _emu(1.1), _emu(3.2), _emu(1.7))
        left.rotation = -12
        left.fill.solid()
        left.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
        _style_shape_text(left, "Rotated shape", 18)

        middle = sld.shapes.add_textbox(_emu(4.7), _emu(0.9), _emu(2.2), _emu(3.1))
        middle.rotation = 90
        p = middle.text_frame.paragraphs[0]
        p.text = "Vertical rotated label"
        p.font.size = Pt(20)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        right = sld.shapes.add_shape(MSO_SHAPE.CAN, _emu(8.2), _emu(1.2), _emu(3.2), _emu(2.3))
        right.rotation = 18
        right.fill.solid()
        right.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
        _style_shape_text(right, "Cylinder", 18)

        conn = sld.shapes.add_connector(MSO_CONNECTOR.CURVE, _emu(3.9), _emu(4.6), _emu(9.8), _emu(4.6))
        conn.line.width = Pt(3)
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        conn.line.color.rgb = RGBColor(0xED, 0x7D, 0x31)
    _add("rotated-text-and-shapes", _build_rotated_text_and_shapes, _patch_connector_tail_arrows)

    # --- Chart + table + shape callout overlay ---
    def _build_chart_table_callout_overlay(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        chart_data = CategoryChartData()
        chart_data.categories = ["North", "South", "East", "West"]
        chart_data.add_series("Actual", (42, 58, 49, 66))
        chart_data.add_series("Plan", (50, 55, 52, 60))
        sld.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            _emu(0.7),
            _emu(0.8),
            _emu(7.1),
            _emu(4.6),
            chart_data,
        )

        tbl_shape = sld.shapes.add_table(3, 2, _emu(8.2), _emu(1.0), _emu(3.9), _emu(1.8))
        tbl = tbl_shape.table
        tbl.cell(0, 0).text = "Metric"
        tbl.cell(0, 1).text = "Value"
        tbl.cell(1, 0).text = "Delta"
        tbl.cell(1, 1).text = "+8%"
        tbl.cell(2, 0).text = "Risk"
        tbl.cell(2, 1).text = "Low"

        callout = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(8.1), _emu(3.25), _emu(4.0), _emu(1.55))
        callout.fill.solid()
        callout.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        callout.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _style_shape_text(callout, "Overlay label should not hide chart/table text", 14)
    _add("chart-table-callout-overlay", _build_chart_table_callout_overlay)

    # --- Translucent overlapping shapes with foreground labels ---
    def _build_layered_transparent_shapes(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        palette = [
            (MSO_SHAPE.OVAL, _emu(1.4), _emu(1.1), RGBColor(0x44, 0x72, 0xC4), "Alpha Layer Blue"),
            (MSO_SHAPE.OVAL, _emu(3.2), _emu(1.1), RGBColor(0xED, 0x7D, 0x31), "Alpha Layer Orange"),
            (MSO_SHAPE.OVAL, _emu(2.3), _emu(2.7), RGBColor(0x70, 0xAD, 0x47), "Alpha Layer Green"),
        ]
        for st, left, top, color, name in palette:
            shp = sld.shapes.add_shape(st, left, top, _emu(3.2), _emu(2.4))
            shp.name = name
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            shp.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        label = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(6.8), _emu(2.0), _emu(4.6), _emu(1.4))
        label.fill.solid()
        label.fill.fore_color.rgb = RGBColor(0x26, 0x26, 0x26)
        _style_shape_text(label, "Alpha overlap + z-order + foreground text", 15)
    _add("layered-transparent-shapes", _build_layered_transparent_shapes, _patch_layered_transparency_case)

    # --- Dense CJK bullet cards with mixed paragraph levels ---
    def _build_dense_cjk_bullet_cards(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cards = [
            ("知识采集", ["多源接入", "权限继承", "增量同步"]),
            ("智能解析", ["版式识别", "表格抽取", "图片理解"]),
            ("精准检索", ["语义召回", "重排优化", "引用追踪"]),
            ("安全应用", ["租户隔离", "审计留痕", "策略管控"]),
        ]
        for idx, (title, bullets) in enumerate(cards):
            row, col = divmod(idx, 2)
            card = sld.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                _emu(0.8 + col * 6.1),
                _emu(0.8 + row * 3.0),
                _emu(5.4),
                _emu(2.3),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFB)
            card.line.color.rgb = RGBColor(0xB4, 0xC7, 0xE7)
            tf = card.text_frame
            tf.word_wrap = True
            p0 = tf.paragraphs[0]
            p0.text = title
            p0.font.name = "Microsoft YaHei"
            p0.font.size = Pt(17)
            p0.font.bold = True
            for bullet in bullets:
                p = tf.add_paragraph()
                p.text = bullet
                p.level = 1
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(12)
    _add("dense-cjk-bullet-cards", _build_dense_cjk_bullet_cards)

    # --- Valid grouped diagram with non-identity child coordinate space ---
    def _build_scaled_group_diagram(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        title = sld.shapes.add_textbox(_emu(0.8), _emu(0.35), _emu(10.5), _emu(0.5))
        p = title.text_frame.paragraphs[0]
        p.text = "Scaled group diagram"
        p.font.size = Pt(24)
        p.font.bold = True
    _add("scaled-group-diagram", _build_scaled_group_diagram, _patch_scaled_group_diagram_case)

    # --- Vertical side label next to a dense table ---
    def _build_vertical_text_with_table(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        side = sld.shapes.add_textbox(_emu(0.45), _emu(1.0), _emu(0.8), _emu(5.2))
        body_pr = side.text_frame._txBody.find(qn("a:bodyPr"))
        body_pr.set("vert", "eaVert")
        p = side.text_frame.paragraphs[0]
        p.text = "纵向标签"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.bold = True

        tbl_shape = sld.shapes.add_table(6, 5, _emu(1.6), _emu(0.9), _emu(10.5), _emu(5.5))
        tbl = tbl_shape.table
        headers = ["能力", "输入", "处理", "输出", "状态"]
        for c, header in enumerate(headers):
            tbl.cell(0, c).text = header
        for r in range(1, 6):
            tbl.cell(r, 0).text = f"模块 {r}"
            tbl.cell(r, 1).text = "PPTX/XML"
            tbl.cell(r, 2).text = "解析 + 渲染"
            tbl.cell(r, 3).text = "HTML/SVG"
            tbl.cell(r, 4).text = "覆盖"
    _add("vertical-text-with-table", _build_vertical_text_with_table)

    # --- Mixed straight, elbow, curve connectors with dash styles ---
    def _build_mixed_dash_connectors(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        coords = [(1.0, 1.0), (5.2, 1.0), (9.2, 1.0), (5.2, 4.5)]
        labels = ["Source", "Transform", "Decision", "Output"]
        for (x, y), label in zip(coords, labels):
            shp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(x), _emu(y), _emu(2.2), _emu(0.9))
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(0xDA, 0xE8, 0xFC)
            shp.line.color.rgb = RGBColor(0x6C, 0x8E, 0xB5)
            _style_shape_text(shp, label, 14, RGBColor(0, 0, 0))
        connectors = [
            (MSO_CONNECTOR.STRAIGHT, 3.2, 1.45, 5.2, 1.45, MSO_LINE_DASH_STYLE.SOLID),
            (MSO_CONNECTOR.ELBOW, 7.4, 1.45, 9.2, 1.45, MSO_LINE_DASH_STYLE.DASH),
            (MSO_CONNECTOR.CURVE, 6.3, 1.9, 6.3, 4.5, MSO_LINE_DASH_STYLE.DASH_DOT),
        ]
        for ctype, x1, y1, x2, y2, dash in connectors:
            conn = sld.shapes.add_connector(ctype, _emu(x1), _emu(y1), _emu(x2), _emu(y2))
            conn.line.width = Pt(2.25)
            conn.line.dash_style = dash
            conn.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    _add("mixed-dash-connectors", _build_mixed_dash_connectors, _patch_connector_tail_arrows)

    # --- Compact report mixing chart, table, connectors, callouts, and CJK text ---
    def _build_mini_report_all_systems(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        title = sld.shapes.add_textbox(_emu(0.5), _emu(0.2), _emu(12.3), _emu(0.55))
        p = title.text_frame.paragraphs[0]
        p.text = "Mini report: text, table, chart, shapes, and connectors"
        p.font.size = Pt(22)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        chart_data = CategoryChartData()
        chart_data.categories = ["T1", "T2", "T3"]
        chart_data.add_series("Pass", (88, 92, 95))
        chart_data.add_series("Fail", (12, 8, 5))
        sld.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, _emu(0.7), _emu(1.0), _emu(5.5), _emu(3.8), chart_data)

        tbl_shape = sld.shapes.add_table(4, 3, _emu(6.7), _emu(1.0), _emu(5.7), _emu(2.4))
        tbl = tbl_shape.table
        for c, header in enumerate(["Area", "Risk", "Owner"]):
            tbl.cell(0, c).text = header
        for r, row in enumerate([("Text", "Low", "A"), ("Chart", "Med", "B"), ("Group", "High", "C")], 1):
            for c, value in enumerate(row):
                tbl.cell(r, c).text = value

        alert = sld.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, _emu(6.9), _emu(4.35), _emu(2.8), _emu(0.85))
        alert.fill.solid()
        alert.fill.fore_color.rgb = RGBColor(0xC0, 0x00, 0x00)
        _style_shape_text(alert, "Review", 15)

        cjk = sld.shapes.add_textbox(_emu(9.9), _emu(4.05), _emu(2.5), _emu(1.35))
        tf = cjk.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "组合场景覆盖"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(18)
        p.font.bold = True

        conn = sld.shapes.add_connector(MSO_CONNECTOR.ELBOW, _emu(6.2), _emu(3.0), _emu(7.2), _emu(4.35))
        conn.line.width = Pt(2)
        conn.line.color.rgb = RGBColor(0x70, 0xAD, 0x47)
    _add("mini-report-all-systems", _build_mini_report_all_systems, _patch_connector_tail_arrows)

    return cases


# ---------------------------------------------------------------------------
# P4: ECharts-renderable 2D charts plus bounded native 3D fallback probes
# ---------------------------------------------------------------------------


def _patch_horizontal_negative_literal_chart_case(pptx_path: Path) -> None:
    """Replace chart references with literals and keep negative inversion explicitly disabled."""
    entries, data_by_name = _read_pptx_entries(pptx_path)
    chart_parts = sorted(name for name in data_by_name if name.startswith("ppt/charts/chart"))
    if len(chart_parts) != 1:
        raise RuntimeError(f"expected one chart part, found {len(chart_parts)}")

    chart_part = chart_parts[0]
    root = etree.fromstring(data_by_name[chart_part])
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    series = root.xpath(".//c:barChart/c:ser", namespaces=ns)
    if len(series) != 1:
        raise RuntimeError(f"expected one bar series, found {len(series)}")

    ser = series[0]
    for container_name, ref_name, cache_name, literal_name in (
        ("cat", "strRef", "strCache", "strLit"),
        ("val", "numRef", "numCache", "numLit"),
    ):
        container = ser.find(qn(f"c:{container_name}"))
        reference = container.find(qn(f"c:{ref_name}")) if container is not None else None
        cache = reference.find(qn(f"c:{cache_name}")) if reference is not None else None
        if container is None or reference is None or cache is None:
            raise RuntimeError(f"missing c:{container_name}/c:{ref_name}/c:{cache_name}")
        literal = etree.Element(qn(f"c:{literal_name}"))
        for child in list(cache):
            cache.remove(child)
            literal.append(child)
        container.replace(reference, literal)

    existing_invert = ser.find(qn("c:invertIfNegative"))
    if existing_invert is not None:
        ser.remove(existing_invert)
    invert = etree.Element(qn("c:invertIfNegative"), val="0")
    category = ser.find(qn("c:cat"))
    ser.insert(ser.index(category) if category is not None else len(ser), invert)

    _replace_pptx_entries(
        pptx_path,
        entries,
        {chart_part: etree.tostring(root, encoding="UTF-8", xml_declaration=True)},
    )


def _patch_3d_chart_case(
    pptx_path: Path,
    *,
    source_tag: str,
    target_tag: str,
    view: tuple[tuple[str, str], ...],
    gap_depth: str | None = None,
) -> None:
    """Convert one python-pptx 2D chart into its schema-compatible 3D variant."""
    entries, data_by_name = _read_pptx_entries(pptx_path)
    chart_parts = sorted(name for name in data_by_name if name.startswith("ppt/charts/chart"))
    if len(chart_parts) != 1:
        raise RuntimeError(f"expected one chart part, found {len(chart_parts)}")

    chart_part = chart_parts[0]
    root = etree.fromstring(data_by_name[chart_part])
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    chart_nodes = root.xpath("./c:chart", namespaces=ns)
    plot_areas = root.xpath(".//c:plotArea", namespaces=ns)
    chart_type_nodes = root.xpath(f".//c:plotArea/c:{source_tag}", namespaces=ns)
    if len(chart_nodes) != 1 or len(plot_areas) != 1 or len(chart_type_nodes) != 1:
        raise RuntimeError(f"expected one c:{source_tag} chart")

    chart = chart_nodes[0]
    plot_area = plot_areas[0]
    chart_type = chart_type_nodes[0]
    chart_type.tag = qn(f"c:{target_tag}")

    if gap_depth is not None:
        gap_depth_node = etree.Element(qn("c:gapDepth"), val=gap_depth)
        axis_ids = chart_type.findall(qn("c:axId"))
        insert_at = chart_type.index(axis_ids[0]) if axis_ids else len(chart_type)
        chart_type.insert(insert_at, gap_depth_node)

    view_3d = etree.Element(qn("c:view3D"))
    for name, value in view:
        etree.SubElement(view_3d, qn(f"c:{name}"), val=value)
    chart.insert(chart.index(plot_area), view_3d)

    _replace_pptx_entries(
        pptx_path,
        entries,
        {chart_part: etree.tostring(root, encoding="UTF-8", xml_declaration=True)},
    )


def _patch_column_3d_chart_case(pptx_path: Path) -> None:
    _patch_3d_chart_case(
        pptx_path,
        source_tag="barChart",
        target_tag="bar3DChart",
        view=(
            ("rotX", "20"),
            ("hPercent", "100"),
            ("rotY", "30"),
            ("depthPercent", "150"),
            ("rAngAx", "1"),
            ("perspective", "30"),
        ),
        gap_depth="150",
    )


def _patch_pie_3d_chart_case(pptx_path: Path) -> None:
    _patch_3d_chart_case(
        pptx_path,
        source_tag="pieChart",
        target_tag="pie3DChart",
        view=(
            ("rotX", "30"),
            ("rotY", "0"),
            ("rAngAx", "0"),
            ("perspective", "30"),
        ),
    )


def _build_chart_cases() -> list[CaseDef]:
    cases: list[CaseDef] = []
    seq = 0

    def _add(slug: str, build_fn, postprocess_fn=None, coverage=None):
        nonlocal seq
        seq += 1
        case = {
            "name": f"oracle-pypptx-chart-{seq:04d}-{slug}",
            "build_fn": build_fn,
        }
        if postprocess_fn is not None:
            case["postprocess_fn"] = postprocess_fn
        if coverage is not None:
            case["coverage"] = coverage
        cases.append(case)

    # --- Column/Bar variants ---
    def _build_col_multi_series(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Q1", "Q2", "Q3", "Q4"]
        cd.add_series("Product A", (45, 52, 48, 61))
        cd.add_series("Product B", (32, 38, 41, 35))
        cd.add_series("Product C", (28, 31, 36, 42))
        sld.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("column-clustered-3series", _build_col_multi_series)

    def _build_col_negative(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
        cd.add_series("Profit/Loss", (15, -8, 22, -12, 5, -3))
        sld.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("column-negative-values", _build_col_negative)

    def _build_col_stacked(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["2021", "2022", "2023", "2024"]
        cd.add_series("Hardware", (120, 135, 142, 158))
        cd.add_series("Software", (85, 102, 118, 131))
        cd.add_series("Services", (45, 52, 68, 79))
        sld.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("column-stacked-3series", _build_col_stacked)

    def _build_col_100_stacked(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["A", "B", "C"]
        cd.add_series("X", (30, 40, 25))
        cd.add_series("Y", (50, 35, 45))
        cd.add_series("Z", (20, 25, 30))
        sld.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("column-100-stacked", _build_col_100_stacked)

    def _build_bar_clustered(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Engineering", "Sales", "Marketing", "Support", "HR"]
        cd.add_series("Headcount", (45, 32, 18, 25, 8))
        sld.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("bar-clustered-single", _build_bar_clustered)

    def _build_bar_stacked(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Dept A", "Dept B", "Dept C"]
        cd.add_series("FY23", (120, 95, 80))
        cd.add_series("FY24", (135, 110, 92))
        sld.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("bar-stacked-2series", _build_bar_stacked)

    # --- Line variants ---
    def _build_line_multi(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
        cd.add_series("Website", (1200, 1350, 1100, 1450, 1380, 1520))
        cd.add_series("Mobile", (800, 920, 850, 1050, 1100, 1180))
        cd.add_series("API", (300, 350, 380, 420, 460, 510))
        sld.shapes.add_chart(XL_CHART_TYPE.LINE, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("line-3series", _build_line_multi)

    def _build_line_markers(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8"]
        cd.add_series("Actual", (82, 85, 79, 91, 88, 94, 87, 96))
        cd.add_series("Target", (85, 85, 85, 85, 90, 90, 90, 90))
        sld.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("line-with-markers", _build_line_markers)

    def _build_line_stacked(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Mon", "Tue", "Wed", "Thu", "Fri"]
        cd.add_series("Email", (120, 132, 101, 134, 90))
        cd.add_series("Chat", (220, 182, 191, 234, 290))
        sld.shapes.add_chart(XL_CHART_TYPE.LINE_STACKED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("line-stacked", _build_line_stacked)

    # --- Pie / Doughnut ---
    def _build_pie_many(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Chrome", "Safari", "Firefox", "Edge", "Other"]
        cd.add_series("Browser Share", (64, 19, 4, 5, 8))
        sld.shapes.add_chart(XL_CHART_TYPE.PIE, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("pie-5-categories", _build_pie_many)

    def _build_pie_exploded(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["A", "B", "C", "D"]
        cd.add_series("Sales", (40, 25, 20, 15))
        sld.shapes.add_chart(XL_CHART_TYPE.PIE_EXPLODED, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("pie-exploded", _build_pie_exploded)

    def _build_doughnut(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Complete", "In Progress", "Not Started"]
        cd.add_series("Status", (65, 20, 15))
        sld.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("doughnut-3-categories", _build_doughnut)

    def _build_doughnut_exploded(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["A", "B", "C"]
        cd.add_series("Values", (50, 30, 20))
        sld.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT_EXPLODED, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("doughnut-exploded", _build_doughnut_exploded)

    # --- Area ---
    def _build_area(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["2020", "2021", "2022", "2023", "2024"]
        cd.add_series("Revenue", (80, 95, 110, 125, 148))
        cd.add_series("Cost", (60, 68, 75, 82, 91))
        sld.shapes.add_chart(XL_CHART_TYPE.AREA, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("area-2series", _build_area)

    def _build_area_stacked(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Q1", "Q2", "Q3", "Q4"]
        cd.add_series("Product", (45, 52, 48, 55))
        cd.add_series("Service", (30, 35, 42, 38))
        cd.add_series("Support", (15, 18, 20, 22))
        sld.shapes.add_chart(XL_CHART_TYPE.AREA_STACKED, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("area-stacked-3series", _build_area_stacked)

    # --- Scatter ---
    def _build_scatter(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = XyChartData()
        s1 = cd.add_series("Cluster A")
        for x, y in [(1.2, 3.1), (2.4, 4.2), (3.1, 2.8), (1.8, 3.6), (2.9, 4.8)]:
            s1.add_data_point(x, y)
        s2 = cd.add_series("Cluster B")
        for x, y in [(5.1, 1.2), (6.3, 2.1), (5.8, 1.8), (7.1, 2.5), (6.0, 0.9)]:
            s2.add_data_point(x, y)
        sld.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("scatter-2-clusters", _build_scatter)

    def _build_scatter_smooth(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = XyChartData()
        s = cd.add_series("Curve")
        for i in range(20):
            x = i * 0.5
            y = math.sin(x) * 3 + 5
            s.add_data_point(x, y)
        sld.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_SMOOTH, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("scatter-smooth-sine", _build_scatter_smooth)

    # --- Radar ---
    def _build_radar(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Speed", "Power", "Agility", "Defense", "Stamina"]
        cd.add_series("Player A", (85, 70, 90, 65, 75))
        cd.add_series("Player B", (70, 85, 65, 80, 90))
        sld.shapes.add_chart(XL_CHART_TYPE.RADAR, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("radar-2series", _build_radar)

    def _build_radar_filled(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = ["Math", "Science", "English", "History", "Art"]
        cd.add_series("Student", (92, 85, 78, 88, 95))
        sld.shapes.add_chart(XL_CHART_TYPE.RADAR_FILLED, _emu(2), _emu(0.5), _emu(8), _emu(6), cd)
    _add("radar-filled-single", _build_radar_filled)

    # --- Bubble ---
    def _build_bubble(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = BubbleChartData()
        s = cd.add_series("Markets")
        s.add_data_point(1.5, 2.5, 10)
        s.add_data_point(3.0, 4.0, 25)
        s.add_data_point(5.0, 1.5, 15)
        s.add_data_point(2.5, 3.5, 30)
        sld.shapes.add_chart(XL_CHART_TYPE.BUBBLE, _emu(1), _emu(0.5), _emu(10), _emu(6), cd)
    _add("bubble-4-points", _build_bubble)

    # --- Large dataset ---
    def _build_line_large(prs):
        sld = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData()
        cd.categories = [str(i) for i in range(1, 25)]
        random.seed(42)
        base = 100
        vals = []
        for _ in range(24):
            base += random.randint(-10, 15)
            vals.append(base)
        cd.add_series("Monthly Trend", vals)
        sld.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, _emu(0.5), _emu(0.5), _emu(11), _emu(6), cd)
    _add("line-24-month-trend", _build_line_large)

    def _build_bar_negative_literal(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Loss", "Gain"]
        data.add_series("Result", (-3, 5))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            _emu(1),
            _emu(0.5),
            _emu(10),
            _emu(6),
            data,
        ).chart
        chart.has_legend = False
        chart.value_axis.minimum_scale = -4
        chart.value_axis.maximum_scale = 6

    _add(
        "bar-negative-literal-zero-crossing",
        _build_bar_negative_literal,
        _patch_horizontal_negative_literal_chart_case,
        coverage={
            "oracle": "native-powerpoint",
            "features": [
                "chart.bar.horizontal",
                "chart.series.strLit",
                "chart.series.numLit",
                "chart.invertIfNegative=false",
                "chart.value-axis.crosses-zero",
                "chart.category-labels.zero-crossing",
            ],
        },
    )

    def _build_column_3d_fallback(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Q1", "Q2", "Q3", "Q4"]
        data.add_series("North", (18, 27, 23, 34))
        data.add_series("South", (14, 22, 30, 26))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            _emu(1),
            _emu(0.5),
            _emu(10),
            _emu(6),
            data,
        )

    _add(
        "column-3d-fallback-view",
        _build_column_3d_fallback,
        _patch_column_3d_chart_case,
        coverage={
            "oracle": "native-powerpoint",
            "features": [
                "chart.bar3DChart",
                "chart.view3D.rotX",
                "chart.view3D.hPercent",
                "chart.view3D.rotY",
                "chart.view3D.depthPercent",
                "chart.view3D.rAngAx",
                "chart.view3D.perspective",
                "chart.fallback.2d-readable",
            ],
        },
    )

    def _build_pie_3d_fallback(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Product", "Services", "Support", "Other"]
        data.add_series("Revenue", (45, 30, 15, 10))
        slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            _emu(2),
            _emu(0.5),
            _emu(8),
            _emu(6),
            data,
        )

    _add(
        "pie-3d-fallback-view",
        _build_pie_3d_fallback,
        _patch_pie_3d_chart_case,
        coverage={
            "oracle": "native-powerpoint",
            "features": [
                "chart.pie3DChart",
                "chart.view3D.rotX",
                "chart.view3D.rotY",
                "chart.view3D.rAngAx",
                "chart.view3D.perspective",
                "chart.fallback.2d-readable",
            ],
        },
    )

    return cases


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------

def _build_all_case_defs(*, include_local_shape3d: bool = False) -> list[CaseDef]:
    all_cases: list[CaseDef] = []
    all_cases.extend(_build_text_cases())
    all_cases.extend(_build_shape_adj_cases())
    all_cases.extend(_build_flowchart_zero_adjustment_cases())
    all_cases.extend(_build_shape_effect_cases())
    all_cases.extend(_build_shape3d_cases())
    if include_local_shape3d:
        all_cases.extend(_build_local_shape3d_cases())
    all_cases.extend(_build_table_cases())
    all_cases.extend(_build_formula_cases())
    all_cases.extend(_build_composite_cases())
    all_cases.extend(_build_chart_cases())
    return all_cases


def _select_case_defs(case_defs: list[CaseDef], patterns: list[str] | None) -> list[CaseDef]:
    """Select case definitions by repeatable exact or shell-style glob patterns."""
    if not patterns:
        return list(case_defs)
    return [
        case_def
        for case_def in case_defs
        if any(fnmatch.fnmatchcase(case_def["name"], pattern) for pattern in patterns)
    ]


def _file_fingerprint(path: Path) -> dict | None:
    if not path.exists() or not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"size_bytes": path.stat().st_size, "sha256": digest.hexdigest()}


def _case_artifact_record(
    name: str,
    status: str,
    pptx_path: Path,
    pdf_path: Path,
    slides_dir: Path,
) -> dict:
    pngs = []
    if slides_dir.is_dir():
        for png_path in sorted(slides_dir.glob("slide*.png")):
            if fingerprint := _file_fingerprint(png_path):
                pngs.append({"name": png_path.name, **fingerprint})
    return {
        "case": name,
        "status": status,
        "source_pptx": _file_fingerprint(pptx_path),
        "ground_truth_pdf": _file_fingerprint(pdf_path),
        "ground_truth_pngs": pngs,
    }


def _generate_pptx(case_def: CaseDef, output_path: str | Path) -> None:
    """Generate a single PPTX file using python-pptx."""
    output_path = Path(output_path)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    case_def["build_fn"](prs)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    if postprocess_fn := case_def.get("postprocess_fn"):
        postprocess_fn(output_path)


def _write_case_json(case_def: CaseDef, cases_dir: Path) -> Path:
    """Write a minimal case JSON for eval script discovery."""
    name = case_def["name"]
    slide_count = int(case_def.get("slide_count", 1))
    payload = {
        "name": name,
        "generator": "python-pptx",
        "slides": [
            {"nodes": [{"kind": "pypptx-generated"}]}
            for _ in range(slide_count)
        ],
    }
    if coverage := case_def.get("coverage"):
        payload["coverage"] = coverage
    if assertions := case_def.get("assertions"):
        payload["assertions"] = assertions
    out = cases_dir / f"{name}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return out


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate ground truth cases using python-pptx + PowerPoint PDF export.",
    )
    parser.add_argument("--cases-dir", type=Path, default=CASES_DIR)
    parser.add_argument(
        "--local-cases-dir",
        type=Path,
        default=LOCAL_SHAPE3D_CASES_DIR,
        help="Ignored metadata directory for opt-in local 3D discovery cases.",
    )
    parser.add_argument("--testdata-dir", type=Path, default=TESTDATA_DIR)
    parser.add_argument("--report-path", type=Path, default=REPORT_PATH)
    parser.add_argument("--pptx-only", action="store_true",
                        help="Only generate PPTX files (skip PDF export). Works on any platform.")
    parser.add_argument("--export-png", action="store_true", default=True,
                        help="Export each slide as PNG (default: enabled, Windows only).")
    parser.add_argument("--no-export-png", action="store_false", dest="export_png",
                        help="Skip PNG export.")
    parser.add_argument("--png-width", type=int, default=0, metavar="PX",
                        help="PNG export width in pixels (0 = PowerPoint default).")
    parser.add_argument("--png-height", type=int, default=0, metavar="PX",
                        help="PNG export height in pixels (0 = PowerPoint default).")
    parser.add_argument("--no-reuse", action="store_true",
                        help="Force regeneration even if files exist.")
    parser.add_argument(
        "--include-local-shape3d-matrix",
        action="store_true",
        help="Include ignored 3D discovery cases without changing supported capability claims.",
    )
    parser.add_argument(
        "--case",
        action="append",
        dest="case_patterns",
        default=[],
        metavar="PATTERN",
        help="Generate only matching case names; repeat for more exact/glob patterns.",
    )
    args = parser.parse_args()

    cases_dir = args.cases_dir.resolve()
    local_cases_dir = args.local_cases_dir.resolve()
    testdata_dir = args.testdata_dir.resolve()
    report_path = args.report_path.resolve()

    all_cases = _build_all_case_defs(include_local_shape3d=args.include_local_shape3d_matrix)
    selected_cases = _select_case_defs(all_cases, args.case_patterns)
    print(f"Total case definitions: {len(all_cases)}; selected: {len(selected_cases)}")
    if args.case_patterns and not selected_cases:
        print(f"ERROR: no cases matched: {args.case_patterns}", file=sys.stderr)
        return 2

    generated: list[str] = []
    failures: list[dict] = []
    artifacts: list[dict] = []
    skipped = 0

    do_png = args.export_png and not args.pptx_only and sys.platform == "win32"

    # Import ground truth export only if needed (per-case native PowerPoint session for stability).
    export_fn = None
    if not args.pptx_only:
        if sys.platform not in {"darwin", "win32"}:
            print(
                "ERROR: PDF export requires macOS or Windows with Microsoft PowerPoint. "
                "Use --pptx-only on other platforms.",
                file=sys.stderr,
            )
            return 1
        from oracle.powerpoint_oracle import export_pptx_ground_truth
        export_fn = export_pptx_ground_truth
        if args.export_png and sys.platform == "darwin":
            print("macOS PowerPoint export is PDF-only; skipping per-slide PNG export.")

    for i, case_def in enumerate(selected_cases, 1):
        name = case_def["name"]
        case_d = testdata_dir / "cases" / name
        pptx_path = case_d / "source.pptx"
        pdf_path = case_d / "ground-truth.pdf"
        slides_d = case_d / "slides"

        # Keep the tracked case index aligned even when ignored binary ground truth is reused.
        definition_dir = local_cases_dir if case_def.get("local_only") else cases_dir
        _write_case_json(case_def, definition_dir)

        # Reuse check
        if not args.no_reuse:
            if args.pptx_only and pptx_path.exists():
                skipped += 1
                artifacts.append(
                    _case_artifact_record(name, "reused", pptx_path, pdf_path, slides_d)
                )
                continue
            if not args.pptx_only and pptx_path.exists() and pdf_path.exists():
                # If PNG export requested but slide1.png missing, regenerate
                if do_png and not (slides_d / "slide1.png").exists():
                    pass  # fall through
                else:
                    skipped += 1
                    artifacts.append(
                        _case_artifact_record(name, "reused", pptx_path, pdf_path, slides_d)
                    )
                    continue

        print(f"  [{i}/{len(selected_cases)}] {name} ...", end=" ", flush=True)

        try:
            # Generate PPTX via python-pptx
            _generate_pptx(case_def, pptx_path)

            # Export PDF + PNG via independent PowerPoint COM session (one per case)
            if export_fn is not None:
                export_fn(
                    pptx_path, pdf_path,
                    slides_png_dir=slides_d if do_png else None,
                    png_width=args.png_width,
                    png_height=args.png_height,
                    runtime_dir=testdata_dir / "oracle-runtime",
                )

            generated.append(name)
            artifacts.append(
                _case_artifact_record(name, "generated", pptx_path, pdf_path, slides_d)
            )
            print("OK")
        except Exception as exc:
            failures.append({"case": name, "error": str(exc)})
            print(f"FAIL: {exc}")

    print(f"\nResults: {len(generated)} generated, {skipped} reused, {len(failures)} failed")

    report = {
        "generator": "python-pptx",
        "schema_version": 2,
        "platform": sys.platform,
        "pptx_only": args.pptx_only,
        "export_png": do_png,
        "png_width": args.png_width,
        "png_height": args.png_height,
        "total_definitions": len(all_cases),
        "selected_definitions": len(selected_cases),
        "case_patterns": args.case_patterns,
        "local_shape3d_matrix": args.include_local_shape3d_matrix,
        "local_cases_dir": str(local_cases_dir),
        "generated_count": len(generated),
        "skipped_reused": skipped,
        "failed_count": len(failures),
        "generated_cases": generated,
        "failed_cases": failures,
        "artifacts": artifacts,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Report: {report_path}")

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
